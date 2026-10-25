from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from docforge.agents.template_analyzer import TemplateAnalyzerAgent
from docforge.design.system import tokens_from_theme
from docforge.llm import FakeLLMClient
from docforge.models.conversation import FileRole, Message, Role, Session, UploadedFile
from docforge.models.intent import Interpretation, OutputFormat, OutputSpec, ResearchNeed
from docforge.models.template import (
    ElementRole,
    LayoutArchetype,
    LayoutGuidance,
    TemplateDesignAnalysis,
    TemplateManifest,
)
from docforge.orchestrator.generate import GenerationService
from docforge.template.pptx_analyzer import analyze_pptx
from docforge.validate.deck import build_deck
from tests.fixtures.decks import sample_deck
from tests.fixtures.templates import SAMPLE_BULLET, SAMPLE_TITLE, corporate_pptx

pytestmark = pytest.mark.anyio


@pytest.fixture
def template(tmp_path: Path) -> Path:
    return corporate_pptx(tmp_path / "corporate.pptx")


def test_analyzer_reads_theme_layouts_and_roles(template: Path):  # FR-TPL-2/3/4
    m = analyze_pptx(template)
    assert m.theme.colors["accent1"] == "C0392B" and m.theme.colors["accent2"] == "2C3E50"
    assert (m.theme.major_font, m.theme.minor_font) == ("Georgia", "Verdana")
    assert m.pptx.slide_width == 9144000 and m.pptx.slide_height == 6858000
    by_name = {lay.name: lay.archetype for lay in m.pptx.layouts}
    assert by_name["Title Slide"] == LayoutArchetype.title
    assert by_name["Title and Content"] == LayoutArchetype.content
    assert by_name["Section Header"] == LayoutArchetype.section
    assert by_name["Title Only"] == LayoutArchetype.title_only
    assert by_name["Comparison"] == LayoutArchetype.comparison
    assert by_name["Blank"] == LayoutArchetype.blank
    assert len(m.pptx.sample_slides) == 3
    roles = {(e.name, e.role) for inv in m.pptx.sample_slides for e in inv.elements}
    assert ("Company Logo", ElementRole.branding) in roles
    assert any(role == ElementRole.sample_content for name, role in roles if name.startswith("Title"))


async def test_template_design_agent_is_evidence_bounded(template: Path):
    manifest = analyze_pptx(template)
    analysis = TemplateDesignAnalysis(
        summary="Observed corporate design.",
        typography={"heading": "Georgia", "invented": "Papyrus"},
        layouts=[
            LayoutGuidance(layout_id=manifest.pptx.layouts[0].id, purpose="Cover"),
            LayoutGuidance(layout_id="invented-layout", purpose="Unsupported"),
        ],
    )
    fake = FakeLLMClient({"template_analyzer": analysis})
    result = await TemplateAnalyzerAgent(fake).run(manifest)

    prompt = fake.calls_for("template_analyzer")[0].messages[0].content
    assert "sample_slides" in prompt and SAMPLE_TITLE not in prompt
    assert [layout.layout_id for layout in result.layouts] == [manifest.pptx.layouts[0].id]
    assert result.typography == {"heading": "Georgia"}


async def test_failed_template_reanalysis_discards_stale_guidance(container, template: Path):
    manifest = analyze_pptx(template).model_copy(
        update={
            "design_analysis": TemplateDesignAnalysis(summary="stale model guidance"),
            "design_analysis_version": 0,
            "design_analysis_model": "fake:old-model",
        }
    )

    refreshed = await GenerationService(container)._analyze_template(manifest)

    assert refreshed.design_analysis is None
    assert refreshed.pptx == manifest.pptx
    assert refreshed.theme == manifest.theme


def test_analyzer_records_explicit_and_inherited_style_provenance(template: Path, tmp_path: Path):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation(str(template))
    slide = prs.slides[1]
    title = slide.shapes.title
    run = title.text_frame.paragraphs[0].runs[0]
    run.font.name = "Georgia"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(192, 57, 43)
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    panel.name = "Explicit Panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(17, 34, 51)
    path = tmp_path / "provenance.pptx"
    prs.save(str(path))

    manifest = analyze_pptx(path)
    inventory = manifest.pptx.sample_slides[1]
    title_element = next(element for element in inventory.elements if element.name == title.name)
    panel_element = next(element for element in inventory.elements if element.name == "Explicit Panel")

    assert title_element.font.name == "Georgia"
    assert title_element.font.name_source.value == "explicit"
    assert title_element.font.size_source.value == "explicit"
    assert title_element.font.bold_source.value == "explicit"
    assert title_element.font.color_source.value == "explicit"
    assert panel_element.fill == "112233" and panel_element.fill_source.value == "explicit"


def test_placeholder_geometry_is_marked_inherited_from_master():
    from docforge.models.template import Box, PropertySource
    from docforge.template.pptx_analyzer import _placeholder_box

    placeholder_type = object()
    layout_placeholder = SimpleNamespace(
        left=None,
        top=None,
        width=None,
        height=None,
        placeholder_format=SimpleNamespace(type=placeholder_type),
    )
    master_placeholder = SimpleNamespace(
        left=100,
        top=200,
        width=300,
        height=400,
        placeholder_format=SimpleNamespace(type=placeholder_type),
    )
    layout = SimpleNamespace(slide_master=SimpleNamespace(placeholders=[master_placeholder]))

    box, source = _placeholder_box(layout_placeholder, layout)

    assert box == Box(x=100, y=200, w=300, h=400)
    assert source == PropertySource.inherited


def test_manifest_without_design_analysis_remains_compatible(template: Path):
    manifest = analyze_pptx(template)
    legacy_data = manifest.model_dump()
    legacy_data.pop("design_analysis")
    legacy_data.pop("design_analysis_version")
    legacy_data.pop("design_analysis_model")

    restored = TemplateManifest.model_validate(legacy_data)

    assert restored.design_analysis is None
    assert restored.design_analysis_version == 0


def test_tokens_follow_theme(template: Path):
    tokens = tokens_from_theme(analyze_pptx(template).theme, "topic")
    assert tokens.source == "template" and tokens.palette.primary == "C0392B"
    assert (tokens.fonts.heading, tokens.fonts.body) == ("Georgia", "Verdana")


def test_pptx_renderer_applies_global_and_slide_scoped_style(tmp_path: Path):
    deck = sample_deck()
    deck.style.heading_font = "Georgia"
    deck.style.body_font = "Verdana"
    deck.style.body_size_pt = 14
    deck.style.palette_overrides = {"text": "112233"}
    target_id = deck.slides[1].id
    deck.slides[1].style.heading_font = "Aptos Display"
    tokens = tokens_from_theme(analyze_pptx(corporate_pptx(tmp_path / "style.pptx")).theme, "style")

    build = build_deck(
        deck,
        tokens,
        tmp_path / "styled-output",
        spec=None,
        dossier=None,
        max_iterations=3,
        libreoffice=None,
    )
    prs = Presentation(str(build.path))

    def named_shape(slide_index: int, name: str):
        return next(shape for shape in prs.slides[slide_index].shapes if shape.name == name)

    first_headline = named_shape(0, "content:title").text_frame.paragraphs[0].runs[0]
    second_headline = named_shape(1, "content:headline").text_frame.paragraphs[0].runs[0]
    body_shape = named_shape(1, "content:item-1")
    body_run = next(run for paragraph in body_shape.text_frame.paragraphs for run in paragraph.runs)
    assert first_headline.font.name == "Georgia"
    assert second_headline.font.name == "Aptos Display"
    assert body_run.font.name == "Verdana" and body_run.font.size.pt == 14
    assert str(body_run.font.color.rgb) == "112233"
    assert deck.slides[1].id == target_id


def test_render_on_template_keeps_identity_and_drops_samples(template: Path, tmp_path: Path):  # FR-PPTX-2/7
    manifest = analyze_pptx(template)
    deck = sample_deck()
    tokens = tokens_from_theme(manifest.theme, deck.title)
    build = build_deck(
        deck,
        tokens,
        tmp_path / "out",
        spec=None,
        dossier=None,
        max_iterations=3,
        libreoffice=None,
        template=template,
        manifest=manifest,
    )
    unresolved = [f.message for f in build.report.findings if f.status != "corrected"]
    assert unresolved == [], unresolved
    checks = {c.check_id: c.status for c in build.report.checks}
    assert (
        checks["pptx.template_fonts"]
        == checks["pptx.sample_residue"]
        == checks["pptx.branding_retained"]
        == "passed"
    )

    prs = Presentation(str(build.path))
    assert len(prs.slides) == len(build.deck.slides)  # 13 new slides, the 3 samples are gone
    assert prs.slide_width == 9144000  # 4:3 template size kept
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == "Title Slide" and layouts[2] == "Section Header" and "Title Only" in layouts
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert SAMPLE_TITLE not in text and SAMPLE_BULLET not in text
    title_ph = prs.slides[3].shapes.title
    assert (
        title_ph is not None and title_ph.name == "content:headline"
    )  # headline in the layout's placeholder
    assert all(any(sh.name.startswith("branding:") for sh in s.shapes) for s in prs.slides)
    assert not any(
        ph.has_text_frame and not ph.text_frame.text.strip() for s in prs.slides for ph in s.placeholders
    )


def test_requested_font_override_is_not_flagged_as_foreign_template_font(template: Path, tmp_path: Path):
    manifest = analyze_pptx(template)
    deck = sample_deck()
    deck.style.heading_font = "Aptos Display"
    build = build_deck(
        deck,
        tokens_from_theme(manifest.theme, deck.title),
        tmp_path / "font-override",
        spec=None,
        dossier=None,
        max_iterations=3,
        libreoffice=None,
        template=template,
        manifest=manifest,
    )
    font_check = next(check for check in build.report.checks if check.check_id == "pptx.template_fonts")
    assert font_check.status == "passed"
    prs = Presentation(str(build.path))
    assert any(
        run.font.name == "Aptos Display"
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    )


def test_template_with_wide_slides(tmp_path: Path):
    template = corporate_pptx(tmp_path / "wide.pptx", widescreen=True)
    manifest = analyze_pptx(template)
    deck = sample_deck()
    build = build_deck(
        deck,
        tokens_from_theme(manifest.theme, "x"),
        tmp_path / "w",
        spec=None,
        dossier=None,
        max_iterations=3,
        libreoffice=None,
        template=template,
        manifest=manifest,
    )
    assert Presentation(str(build.path)).slide_width == manifest.pptx.slide_width
    assert [f.message for f in build.report.findings if f.status != "corrected"] == []


async def test_generation_uses_uploaded_template_and_caches_manifest(container, fake_llm, template: Path):
    ref = container.blobs.put_file(template)
    container.repos.sessions.put(Session(id="ses_1", owner="demo"))
    container.repos.files.put(
        UploadedFile(
            id="fil_t",
            session_id="ses_1",
            filename="corporate.pptx",
            media_type="x",
            size=ref.size,
            sha256=ref.sha256,
            role=FileRole.template,
            blob_path=ref.path,
        )
    )
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind="new_task",
            confidence=0.9,
            objective="EV deck",
            output_spec=OutputSpec(format=OutputFormat.pptx),
            research=ResearchNeed(depth="none"),
        ),
    )
    fake_llm.add(
        "template_analyzer",
        TemplateDesignAnalysis(
            summary="Use the observed Georgia/Verdana typography and preserve the logo.",
            typography={"heading": "Georgia", "body": "Verdana"},
            color_roles={"accent": "accent1"},
            preservation_rules=["Keep the company logo and template masters."],
        ),
    )
    fake_llm.add("deck_planner", sample_deck())
    session = container.repos.sessions.require("ses_1")
    msg = container.repos.messages.put(Message(id="m1", session_id="ses_1", role=Role.user, content="deck"))
    reply = await GenerationService(container).handle_message(session, msg)

    assert "Built on your template “corporate.pptx”" in reply.message.content
    assert "Template 'corporate.pptx'" in fake_llm.calls_for("deck_planner")[0].messages[0].content
    assert container.repos.templates.get(ref.sha256) is not None
    cached = container.repos.templates.require(ref.sha256)
    assert cached.design_analysis is not None
    assert cached.design_analysis.summary.startswith("Use the observed")
    assert len(fake_llm.calls_for("template_analyzer")) == 1
    assert "Template design analysis" in fake_llm.calls_for("deck_planner")[0].messages[0].content
    assert reply.version.template_file_id == "fil_t"
    assert container.repos.tasks.require(reply.task.id).template_file_id == "fil_t"

    container.settings.llm_agent_model = {"template_analyzer": "fake:updated-model"}
    refreshed = await GenerationService(container)._analyze_template(cached)
    assert refreshed.design_analysis_model == "fake:updated-model"
    assert len(fake_llm.calls_for("template_analyzer")) == 2
    prs = Presentation(io.BytesIO(container.blobs.read_bytes(reply.version.blob_path)))
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert reply.version.design_tokens.source == "template"


def test_active_deck_uses_its_pinned_template(container, tmp_path: Path):
    from docforge.models.conversation import Task, TaskStatus

    session = Session(id="ses_pin", owner="demo")
    container.repos.sessions.put(session)
    refs = []
    for file_id, filename in (("fil_old", "old.pptx"), ("fil_new", "new.pptx")):
        path = corporate_pptx(tmp_path / filename)
        ref = container.blobs.put_file(path)
        refs.append(ref)
        container.repos.files.put(
            UploadedFile(
                id=file_id,
                session_id=session.id,
                filename=filename,
                media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                size=ref.size,
                sha256=ref.sha256,
                role=FileRole.template,
                blob_path=ref.path,
            )
        )
    task = Task(
        id="tsk_pin",
        session_id=session.id,
        objective="presentation",
        status=TaskStatus.delivered,
        template_file_id="fil_old",
    )
    session.active_task_id = task.id
    container.repos.sessions.put(session)
    container.repos.tasks.put(task)

    service = GenerationService(container)
    selected = service._template_path(session, "pptx")

    assert selected == container.blobs.resolve(refs[0].path)
    service._referenced = {"fil_new"}
    selected_replacement = service._template_path(session, "pptx")
    assert selected_replacement == container.blobs.resolve(refs[1].path)
    assert container.repos.tasks.require(task.id).template_file_id == "fil_new"


def test_analyze_docx_handles_styles_without_base_style(tmp_path: Path):
    from docx import Document

    from docforge.template.docx_analyzer import analyze_docx

    doc_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_heading("Test Document", level=0)
    doc.add_paragraph("Sample body text")
    doc.save(str(doc_path))

    manifest = analyze_docx(doc_path)
    assert manifest.docx is not None
    assert len(manifest.docx.styles) > 0
    # Also verify with a style mock that has neither base_style nor font
    from docforge.template.docx_analyzer import _font

    mock_style = SimpleNamespace(name="No List", style_id="NoList")
    assert _font(mock_style) is None
