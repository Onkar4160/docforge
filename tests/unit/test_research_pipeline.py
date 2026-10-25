from __future__ import annotations

import io
import json

import httpx
import pytest
from docx import Document
from pptx import Presentation

from docforge.agents.research import ExtractedConflict, ExtractedItem, Extraction, PlannedResearch
from docforge.core.errors import ResearchError
from docforge.ingest.readers import extract_text
from docforge.llm import FakeLLMClient
from docforge.models.evidence import Authority, ResearchQuestion
from docforge.models.intent import Interpretation, OutputFormat, OutputSpec, ResearchDepth, ResearchNeed
from docforge.models.ir import (
    BulletItem,
    BulletsBlock,
    ChartBlock,
    ChartSeries,
    DeckPlan,
    KPIBlock,
    KPIMetric,
    SlidePlan,
)
from docforge.orchestrator.generate import GenerationService
from docforge.research.academic import (
    ArxivProvider,
    OpenAlexProvider,
    SemanticScholarProvider,
    _abstract_from_index,
    arxiv_query,
)
from docforge.research.pipeline import ResearchService, UserSource, authority_of, evidence_digest, rank_score
from docforge.research.search import FixtureSearchProvider, SearchQuery, SearchResult
from docforge.validate.common import CheckRun
from docforge.validate.plan_checks import check_deck_plan
from tests.conftest import make_settings

pytestmark = pytest.mark.anyio


def _client(handler) -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


# ---------------------------------------------------------------- scholarly providers


def test_openalex_abstract_reconstruction():
    assert _abstract_from_index({"charging": [1], "Fast": [0], "wins": [2]}) == "Fast charging wins"
    assert _abstract_from_index(None) == ""


async def test_openalex_request_and_mapping():
    seen: list[httpx.Request] = []

    def handler(req):
        seen.append(req)
        return httpx.Response(
            200,
            json={
                "results": [
                    {
                        "id": "https://openalex.org/W1",
                        "doi": "https://doi.org/10.1/abc",
                        "display_name": "EV study",
                        "publication_year": 2025,
                        "publication_date": "2025-03-01",
                        "cited_by_count": 42,
                        "type": "article",
                        "primary_location": {"source": {"display_name": "Nature Energy"}},
                        "authorships": [{"author": {"display_name": "A. Author"}}],
                        "abstract_inverted_index": {"Charging": [0], "grows": [1]},
                    }
                ]
            },
        )

    results = await OpenAlexProvider("oa-key", mailto="me@x.org", client=_client(handler)).search(
        SearchQuery("ev charging", max_results=5)
    )
    r = results[0]
    assert (r.title, r.url, r.content, r.venue, r.citations, r.kind) == (
        "EV study",
        "https://doi.org/10.1/abc",
        "Charging grows",
        "Nature Energy",
        42,
        "paper",
    )
    params = dict(seen[0].url.params)
    assert params["api_key"] == "oa-key" and params["mailto"] == "me@x.org" and params["per_page"] == "5"


async def test_semantic_scholar_mapping_and_rate_limit():
    def ok(req):
        assert req.headers["x-api-key"] == "s2"
        return httpx.Response(
            200,
            json={
                "data": [
                    {
                        "title": "Grid study",
                        "abstract": None,
                        "tldr": {"text": "Grids constrain growth."},
                        "year": 2024,
                        "venue": "Joule",
                        "authors": [{"name": "B"}],
                        "citationCount": 7,
                        "url": "https://s2/p1",
                        "externalIds": {"DOI": "10.2/x"},
                    },
                    {"title": "No abstract", "abstract": None, "url": "https://s2/p2"},
                ]
            },
        )

    results = await SemanticScholarProvider("s2", client=_client(ok)).search(SearchQuery("grid"))
    assert len(results) == 1 and results[0].content == "Grids constrain growth."
    assert results[0].url == "https://doi.org/10.2/x"
    with pytest.raises(ResearchError, match="rate limit"):
        await SemanticScholarProvider(client=_client(lambda r: httpx.Response(429))).search(SearchQuery("x"))


async def test_arxiv_atom_parsing():
    atom = """<?xml version="1.0"?><feed xmlns="http://www.w3.org/2005/Atom"
      xmlns:arxiv="http://arxiv.org/schemas/atom"><entry><id>http://arxiv.org/abs/2501.00001v1</id>
      <title>  Smart   charging </title><summary> We study smart charging. </summary>
      <published>2025-01-02T00:00:00Z</published><author><name>C. Doe</name></author></entry></feed>"""
    results = await ArxivProvider(client=_client(lambda r: httpx.Response(200, text=atom))).search(
        SearchQuery("smart charging")
    )
    assert (results[0].title, results[0].kind, results[0].published_date) == (
        "Smart charging",
        "preprint",
        "2025-01-02",
    )
    assert arxiv_query("What drives EV charging adoption?") == "all:drives AND all:charging AND all:adoption"


def test_authority_and_ranking():
    paper = SearchResult("p", "https://doi.org/1", content="x", kind="paper", venue="Lancet", citations=500)
    gov = SearchResult("g", "https://www.eia.gov/report", snippet="x")
    blog = SearchResult("b", "https://someone.medium.com/post", snippet="x")
    assert authority_of(paper)[0] == Authority.high and authority_of(gov)[0] == Authority.high
    assert authority_of(blog)[0] == Authority.low
    assert rank_score(paper, 0) > rank_score(blog, 0)


# ---------------------------------------------------------------- ingestion


def test_extract_text_from_office_files():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Revenue grew 12% in 2025.")
    doc.save(buf)
    assert "Revenue grew 12%" in extract_text("r.docx", buf.getvalue())
    buf = io.BytesIO()
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Old deck"
    prs.save(buf)
    assert "[slide 1]" in extract_text("d.pptx", buf.getvalue())
    assert extract_text("t.csv", b"a,b\n1,2\n") == "a\tb\n1\t2"


# ---------------------------------------------------------------- pipeline


def _plan() -> PlannedResearch:
    return PlannedResearch(
        questions=[
            ResearchQuestion(
                id="x",
                question="How big is the market?",
                priority=1,
                query_hints=["ev market size"],
                channels=["web", "academic"],
            ),
            ResearchQuestion(
                id="y",
                question="What limits growth?",
                priority=2,
                query_hints=["grid limits"],
                channels=["academic"],
            ),
            ResearchQuestion(id="z", question="Unanswerable question", priority=3, query_hints=["nothing"]),
        ]
    )


def _web() -> FixtureSearchProvider:
    return FixtureSearchProvider(
        {
            "ev market": [
                SearchResult(
                    "IEA outlook",
                    "https://www.iea.org/outlook",
                    snippet="s",
                    content="Market reached 9.4 bn " * 10,
                    score=0.9,
                ),
                SearchResult(
                    "IEA outlook copy",
                    "https://iea.org/outlook/",
                    snippet="s",
                    content="dup " * 40,
                    score=0.5,
                ),
            ]
        }
    )


def _academic() -> FixtureSearchProvider:
    p = FixtureSearchProvider(
        {
            "grid": [
                SearchResult(
                    "Grid constraints",
                    "https://doi.org/10.9/g",
                    content="Grid connection delays of 18-36 months limit rollout. " * 3,
                    kind="paper",
                    venue="Energy Policy",
                    doi="10.9/g",
                )
            ]
        }
    )
    p.name = "openalex"
    return p


def _extraction(req) -> Extraction:
    text = req.messages[-1].content
    items = []
    if 'id="S1"' in text:
        items.append(
            ExtractedItem(
                source_id="S1",
                kind="figure",
                statement="The market reached EUR 9.4bn in 2025.",
                value=9.4,
                unit="EUR bn",
                as_of="2025",
                question_ids=["q1"],
                confidence=0.9,
            )
        )
        items.append(
            ExtractedItem(
                source_id="S1",
                kind="fact",
                statement="Fast charging leads growth.",
                question_ids=["q1"],
                confidence=0.8,
            )
        )
    if 'id="S2"' in text:
        items.append(
            ExtractedItem(
                source_id="S2",
                kind="finding",
                statement="Grid delays of 18-36 months.",
                question_ids=["q2"],
                confidence=0.85,
            )
        )
    if 'id="U1"' in text:
        items.append(
            ExtractedItem(
                source_id="U1",
                kind="fact",
                statement="Our pilot had 14% utilisation.",
                question_ids=["q1"],
                confidence=0.9,
            )
        )
    items.append(ExtractedItem(source_id="S99", kind="fact", statement="hallucinated source"))
    return Extraction(
        items=items, conflicts=[ExtractedConflict(topic="size", source_ids=["S1"], description="d")]
    )


async def test_research_pipeline_end_to_end(tmp_path):
    fake = FakeLLMClient({"research_planner": _plan(), "evidence_extractor": _extraction})
    svc = ResearchService(make_settings(tmp_path), fake, web=_web(), academic=[_academic()], knowledge=None)
    dossier = await svc.run(
        "tsk_1",
        "EV charging market",
        deliverable="pptx",
        depth=ResearchDepth.standard,
        namespace="ses_1",
        user_sources=[UserSource("fil_1", "pilot.txt", "Pilot data " * 20)],
    )
    assert dossier.performed and [q.id for q in dossier.plan.questions] == ["q1", "q2", "q3"]
    assert {s.id for s in dossier.sources} == {"S1", "S2", "U1"}  # duplicate URL collapsed, unused dropped
    assert len(dossier.claims) == 4 and all(c.evidence_ids for c in dossier.claims)  # S99 discarded
    assert (
        dossier.coverage.answered == 1 and dossier.coverage.partial == 1 and dossier.coverage.unanswered == 1
    )
    assert any("Unanswerable" in g.description for g in dossier.gaps)
    assert dossier.sources[0].venue == "Energy Policy"  # peer-reviewed paper outranks the web page
    assert all(s.authority == Authority.high for s in dossier.sources)

    digest = evidence_digest(dossier)
    assert "[c1] The market reached EUR 9.4bn in 2025. [9.4 EUR bn, 2025]" in digest
    assert "Gaps" in digest and "Conflicting evidence" in digest


async def test_research_records_provider_failures(tmp_path):
    class Broken(FixtureSearchProvider):
        name = "openalex"

        async def search(self, q):
            raise ResearchError("openalex rate limit reached")

    fake = FakeLLMClient({"research_planner": _plan(), "evidence_extractor": _extraction})
    svc = ResearchService(make_settings(tmp_path), fake, web=None, academic=[Broken({})])
    dossier = await svc.run("t", "o", deliverable="pptx", depth=ResearchDepth.light, namespace="n")
    assert any("rate limit" in n for n in dossier.notes) and "no usable sources" in dossier.notes[-1]
    assert dossier.coverage.unanswered == 3


# ---------------------------------------------------------------- citation enforcement + generation


def test_uncited_figures_flagged_when_research_exists(tmp_path):
    from docforge.models.evidence import Claim, EvidenceItem, ResearchDossier, Source

    dossier = ResearchDossier(
        id="d",
        task_id="t",
        performed=True,
        sources=[Source(id="S1", kind="web", title="t", url="https://x.org")],
        evidence=[EvidenceItem(id="e1", source_id="S1", kind="fact", statement="s")],
        claims=[Claim(id="c1", text="s", evidence_ids=["e1"], status="supported")],
    )
    deck = DeckPlan(
        title="t",
        slides=[
            SlidePlan(
                id="s01",
                purpose="p",
                archetype="bullets",
                headline="h",
                blocks=[
                    BulletsBlock(
                        items=[
                            BulletItem(text="Market grew 21% last year"),
                            BulletItem(text="cited 9%", claim_ids=["c1"]),
                            BulletItem(text="qualitative point"),
                        ]
                    ),
                    KPIBlock(metrics=[KPIMetric(label="Size", value="EUR 9bn")]),
                    ChartBlock(
                        chart_type="bar",
                        categories=["a"],
                        series=[ChartSeries(name="s", values=[1])],
                        claim_ids=["c1"],
                    ),
                ],
            )
        ],
    )
    run = CheckRun()
    check_deck_plan(run, deck, None, dossier)
    messages = [f.message for f in run.findings if f.check_id == "evidence.coverage"]
    assert len(messages) == 2 and "21%" in messages[0] and "EUR 9bn" in messages[1]


async def test_generation_with_research_cites_sources(container, fake_llm):
    from docforge.models.conversation import Message, Role, Session
    from tests.fixtures.decks import sample_deck

    container.search = _web()
    container.academic = [_academic()]
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind="new_task",
            confidence=0.9,
            objective="EV charging board update",
            output_spec=OutputSpec(format=OutputFormat.pptx),
            research=ResearchNeed(depth="standard"),
        ),
    )
    fake_llm.add("research_planner", _plan())
    fake_llm.add("evidence_extractor", _extraction)
    deck = sample_deck()
    deck.slides[3].blocks[0].claim_ids = ["c1"]  # chart cites the market figure
    deck.slides[4].blocks[0].claim_ids = ["c1", "c3"]
    deck.slides[8].blocks[0].claim_ids = ["c2"]
    fake_llm.add("deck_planner", deck)

    container.settings.research_enabled = True
    session = container.repos.sessions.put(Session(id="ses_1", owner="demo"))
    msg = container.repos.messages.put(Message(id="m1", session_id="ses_1", role=Role.user, content="deck"))
    reply = await GenerationService(container).handle_message(session, msg)

    assert reply.dossier is not None and reply.dossier.performed
    assert "Research: 2 sources used (1 scholarly)" in reply.message.content
    assert "<evidence>" in fake_llm.calls_for("deck_planner")[0].messages[0].content
    assert "[c1]" in fake_llm.calls_for("deck_planner")[0].messages[0].content
    prs = Presentation(io.BytesIO(container.blobs.read_bytes(reply.version.blob_path)))
    footers = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.name == "footer:sources"]
    assert "Sources: [1]" in footers
    last = " ".join(sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame)
    assert "IEA outlook" in last and "[1]" in last  # auto-appended references slide
    assert container.repos.dossiers.require(reply.dossier.id).task_id == reply.task.id
    assert json.loads(reply.version.model_dump_json())["dossier_id"] == reply.dossier.id
