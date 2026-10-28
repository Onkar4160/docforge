from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from docforge.agents.doc_writer import DocumentInput, assemble, target_words, write_document
from docforge.design.system import generated_tokens
from docforge.llm import FakeLLMClient
from docforge.models.conversation import FileRole, Message, Role, Session, UploadedFile
from docforge.models.evidence import Claim, EvidenceItem, ResearchDossier, Source
from docforge.models.intent import (
    ChangeOperation,
    ChangeRequest,
    ChangeScope,
    CitationStyle,
    ConfirmationDecision,
    Interpretation,
    LengthSpec,
    OutputFormat,
    OutputSpec,
    RequestKind,
    ResearchNeed,
    ScopeLevel,
    SlidePlacement,
    SlidePlacementMode,
    StyleChange,
    StylePlan,
    TargetKind,
    TargetRef,
)
from docforge.models.ir import ParagraphBlock, SectionDraft, SectionInsertionPlan
from docforge.models.quality import Dimension, Severity, Verdict
from docforge.orchestrator.generate import GenerationService
from docforge.render.citations import Citations
from docforge.render.docx import render_document
from docforge.template.docx_analyzer import analyze_docx
from docforge.validate.common import CheckRun
from docforge.validate.document import build_document, check_docx_file
from tests.fixtures.documents import outline, written

pytestmark = pytest.mark.anyio


def _dossier() -> ResearchDossier:
    return ResearchDossier(
        id="d",
        task_id="t",
        performed=True,
        sources=[
            Source(
                id="S1",
                kind="web",
                title="Microplastics in drinking-water",
                url="https://who.int/mp",
                publisher="WHO",
                published_at="2019-08-22",
            ),
            Source(
                id="S2",
                kind="web",
                title="Global review",
                url="https://doi.org/10.1/x",
                author="Koelmans, Mohamed",
                published_at="2019",
                venue="Water Research",
                doi="10.1/x",
            ),
        ],
        evidence=[
            EvidenceItem(id="e1", source_id="S1", kind="fact", statement="a"),
            EvidenceItem(id="e2", source_id="S2", kind="finding", statement="b"),
        ],
        claims=[
            Claim(id="c1", text="a", evidence_ids=["e1"], status="supported"),
            Claim(id="c2", text="b", evidence_ids=["e2"], status="supported"),
        ],
    )


def _plan(cite=None):
    return assemble(outline(), written(cite))


def test_assemble_assigns_ids_and_levels():
    plan = _plan()
    assert [s.id for s in plan.all_sections()] == ["sec-1", "sec-2", "sec-2-1", "sec-2-2"]
    assert plan.sections[1].subsections[0].level == 2
    assert plan.front_matter.executive_summary[0].text.startswith("Microplastics are widespread")


def test_target_words():
    assert (
        target_words(OutputSpec(format="docx", length=LengthSpec(unit="pages", target=10, explicit=True)))
        == 4500
    )
    assert target_words(OutputSpec(format="docx")) == 2800


def test_citation_styles():
    d = _dossier()
    numeric = Citations(d, CitationStyle.numeric)
    assert numeric.marker(["c2", "c1"]) == " [1, 2]" and numeric.marker(["c1"]) == " [2]"
    assert numeric.reference_entries()[0].startswith("[1] Koelmans, Mohamed")
    apa = Citations(d, CitationStyle.apa)
    assert apa.marker(["c2"]) == " (Koelmans et al., 2019)" and apa.marker(["c1"]) == " (WHO, 2019)"
    assert apa.reference_entries()[0].startswith("Koelmans, Mohamed (2019). Global review. Water Research.")
    assert Citations(d, CitationStyle.harvard).marker(["c1"]) == " (WHO 2019)"
    assert Citations(None, CitationStyle.numeric).marker(["c1"]) == ""


def test_render_document_structure(tmp_path: Path):
    plan = _plan(["c1", "c2"])
    result = render_document(
        plan, generated_tokens("water policy"), tmp_path / "r.docx", dossier=_dossier(), audience="Board"
    )
    doc = Document(str(tmp_path / "r.docx"))
    styles = [p.style.name for p in doc.paragraphs]
    assert styles.count("Title") == 1 and "List Bullet" in styles and "List Number" in styles
    assert [h for h in result.headings][:2] == [(1, "Executive summary"), (1, "What the evidence shows")]
    assert (2, "Monitoring") in result.headings and result.headings[-1] == (1, "References")
    assert result.figures == 1 and result.tables == 3 and result.references == 2  # data, timeline, comparison
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "[1]" in text and "Figure 1: Particles per litre by study (particles/L)" in text
    assert "Right-click and choose Update Field" not in text
    assert 'TOC \\o "1-3"' in doc.element.body.xml
    assert "PAGE" in doc.sections[0].footer._element.xml
    assert doc.inline_shapes[0]._inline.docPr.get("descr") == "Particles per litre by study"


def test_template_styles_are_reused(tmp_path: Path):
    template = Document()
    template.styles["Heading 1"].font.name = "Georgia"
    template.add_paragraph("SAMPLE CONTENT TO BE REMOVED")
    template.sections[0].header.paragraphs[0].text = "ACME Corp confidential"
    template_path = tmp_path / "brand.docx"
    template.save(str(template_path))

    render_document(_plan(), generated_tokens("x"), tmp_path / "out.docx", template=template_path)
    doc = Document(str(tmp_path / "out.docx"))
    assert doc.styles["Heading 1"].font.name == "Georgia"
    assert "SAMPLE CONTENT" not in "\n".join(p.text for p in doc.paragraphs)
    assert doc.sections[0].header.paragraphs[0].text == "ACME Corp confidential"


def test_build_document_checks(tmp_path: Path):
    spec = OutputSpec(format="docx", length=LengthSpec(unit="pages", target=40, explicit=True))
    build = build_document(
        _plan(), generated_tokens("x"), tmp_path, spec=spec, dossier=_dossier(), libreoffice=None
    )
    found = {f.check_id for f in build.report.findings}
    assert "common.length" in found  # far shorter than 40 pages
    assert "evidence.coverage" in found  # chart/table/kpi without citations while research exists
    assert build.report.verdict != Verdict.passed
    clean = build_document(
        _plan(["c1", "c2"]),
        generated_tokens("x"),
        tmp_path / "b",
        spec=OutputSpec(format="docx"),
        dossier=_dossier(),
        libreoffice=None,
    )
    assert not [f for f in clean.report.findings if f.severity in ("critical", "major")], (
        clean.report.findings
    )


async def test_write_document_runs_outline_then_parallel_sections():
    sections = written()
    fake = FakeLLMClient({"doc_outliner": outline(), "section_writer": sections})
    plan = await write_document(fake, DocumentInput(objective="o", spec=OutputSpec(format="docx")))
    assert len(fake.calls_for("section_writer")) == 2
    prompt = fake.calls_for("section_writer")[1].messages[0].content
    assert "Write exactly these subsections" in prompt and "Monitoring" in prompt
    assert len(plan.sections) == 2


async def test_generation_service_produces_docx(container, fake_llm):
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind="new_task",
            confidence=0.9,
            objective="Microplastics report",
            output_spec=OutputSpec(format=OutputFormat.docx, audience="board"),
            research=ResearchNeed(depth="none"),
        ),
    )
    fake_llm.add("doc_outliner", outline())
    fake_llm.add("section_writer", *written())
    session = container.repos.sessions.put(Session(id="ses_1", owner="demo"))
    msg = container.repos.messages.put(Message(id="m1", session_id="ses_1", role=Role.user, content="report"))
    reply = await GenerationService(container).handle_message(session, msg)
    assert reply.artifact.format == OutputFormat.docx and reply.version.version == 1
    assert "Created a Word document" in reply.message.content
    doc = Document(io.BytesIO(container.blobs.read_bytes(reply.version.blob_path)))
    assert any(p.text == "Options for the authority" for p in doc.paragraphs)


async def test_generation_service_edits_docx(container, fake_llm):
    # 1. Initial document creation (version 1)
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind="new_task",
            confidence=0.9,
            objective="Microplastics report",
            output_spec=OutputSpec(format=OutputFormat.docx, audience="board"),
            research=ResearchNeed(depth="none"),
        ),
    )
    fake_llm.add("doc_outliner", outline())
    fake_llm.add("section_writer", *written())
    session = container.repos.sessions.put(Session(id="ses_edit", owner="demo"))
    msg1 = container.repos.messages.put(
        Message(id="m1", session_id="ses_edit", role=Role.user, content="create report")
    )
    reply1 = await GenerationService(container).handle_message(session, msg1)
    assert reply1.version.version == 1

    # 2. Dynamic restyle (version 2)
    fake_llm._queues["interpreter"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            confirmation=ConfirmationDecision.approve,
            change=ChangeRequest(
                operation=ChangeOperation.restyle,
                scope=ChangeScope(level=ScopeLevel.style),
                instruction="Use Georgia headings and navy primary color",
                style_changes=[
                    StyleChange(property="heading_font", value="Georgia"),
                    StyleChange(property="primary", value="#000080"),
                ],
            ),
        ),
    )
    fake_llm.add(
        "style_planner",
        StylePlan(
            changes=[
                StyleChange(property="heading_font", value="Georgia"),
                StyleChange(property="primary", value="#000080"),
            ]
        ),
    )
    msg2 = container.repos.messages.put(
        Message(
            id="m2",
            session_id="ses_edit",
            role=Role.user,
            content="Use Georgia headings and navy primary color",
        )
    )
    reply2 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg2
    )
    assert reply2.version.version == 2
    assert reply2.version.ir.style.heading_font == "Georgia"
    assert reply2.version.ir.style.palette_overrides["primary"] == "000080"
    doc2 = Document(io.BytesIO(container.blobs.read_bytes(reply2.version.blob_path)))
    assert doc2.styles["Heading 1"].font.name == "Georgia"

    # 3. Add section dynamically (version 3)
    fake_llm._queues["interpreter"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            change=ChangeRequest(
                operation=ChangeOperation.add,
                scope=ChangeScope(level=ScopeLevel.artifact),
                instruction="Add an implementation roadmap section at the end",
                placement=SlidePlacement(mode=SlidePlacementMode.end),
            ),
        ),
    )
    draft = SectionDraft(
        heading="Implementation Roadmap & Next Steps",
        purpose="action plan",
        level=1,
        blocks=[ParagraphBlock(text="Step 1: Water quality pilot testing.")],
    )
    fake_llm.add("section_inserter", SectionInsertionPlan(sections=[draft]))
    msg3 = container.repos.messages.put(
        Message(
            id="m3",
            session_id="ses_edit",
            role=Role.user,
            content="Add an implementation roadmap section at the end",
        )
    )
    reply3 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg3
    )
    assert reply3.version.version == 3
    assert len(reply3.version.ir.sections) == 3
    assert reply3.version.ir.sections[-1].heading == "Implementation Roadmap & Next Steps"
    doc3 = Document(io.BytesIO(container.blobs.read_bytes(reply3.version.blob_path)))
    assert any("Implementation Roadmap & Next Steps" in p.text for p in doc3.paragraphs)

    # 4. Remove section dynamically (version 4)
    fake_llm._queues["interpreter"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            change=ChangeRequest(
                operation=ChangeOperation.remove,
                scope=ChangeScope(
                    level=ScopeLevel.unit,
                    targets=[TargetRef(kind=TargetKind.section, locator="second section")],
                ),
                instruction="Remove the second section",
            ),
        ),
    )
    msg4 = container.repos.messages.put(
        Message(id="m4", session_id="ses_edit", role=Role.user, content="Remove the second section")
    )
    reply4 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg4
    )
    assert reply4.version.version == 4
    assert len(reply4.version.ir.sections) == 2

    # 5. Edit section content dynamically (version 5)
    target_sec = reply4.version.ir.sections[0]
    edited_sec = target_sec.model_copy(update={"heading": "Empirical Findings & Data"})
    fake_llm._queues["interpreter"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            change=ChangeRequest(
                operation=ChangeOperation.rewrite,
                scope=ChangeScope(
                    level=ScopeLevel.unit,
                    targets=[TargetRef(kind=TargetKind.section, locator="first section")],
                ),
                instruction="Rename first section to Empirical Findings & Data",
            ),
        ),
    )
    fake_llm.add("section_editor", edited_sec)
    msg5 = container.repos.messages.put(
        Message(
            id="m5",
            session_id="ses_edit",
            role=Role.user,
            content="Rename first section to Empirical Findings & Data",
        )
    )
    reply5 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg5
    )
    assert reply5.version.version == 5
    assert reply5.version.ir.sections[0].heading == "Empirical Findings & Data"
    doc5 = Document(io.BytesIO(container.blobs.read_bytes(reply5.version.blob_path)))
    assert any("Empirical Findings & Data" in p.text for p in doc5.paragraphs)

    # 6. Edit section content by adding paragraph (version 6)
    sec6 = reply5.version.ir.sections[0]
    new_block = ParagraphBlock(text="Global monitoring initiatives have expanded empirical testing.")
    sec6_edited = sec6.model_copy(update={"blocks": [*sec6.blocks, new_block]})
    fake_llm._queues["interpreter"].clear()
    fake_llm._queues["section_editor"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            change=ChangeRequest(
                operation=ChangeOperation.rewrite,
                scope=ChangeScope(
                    level=ScopeLevel.unit,
                    targets=[TargetRef(kind=TargetKind.section, locator="first section")],
                ),
                instruction="In the first section, add a paragraph about global monitoring",
            ),
        ),
    )
    fake_llm.add("section_editor", sec6_edited)
    msg6 = container.repos.messages.put(
        Message(
            id="m6",
            session_id="ses_edit",
            role=Role.user,
            content="In the first section, add a paragraph about global monitoring",
        )
    )
    reply6 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg6
    )
    assert reply6.version.version == 6
    doc6 = Document(io.BytesIO(container.blobs.read_bytes(reply6.version.blob_path)))
    assert any("Global monitoring initiatives" in p.text for p in doc6.paragraphs)

    # 7. Exact string replacement in section (version 7) without LLM editor call
    fake_llm._queues["interpreter"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            change=ChangeRequest(
                operation=ChangeOperation.rewrite,
                scope=ChangeScope(
                    level=ScopeLevel.unit,
                    targets=[TargetRef(kind=TargetKind.section, locator="first section")],
                ),
                instruction=(
                    'In the first section, replace "Global monitoring initiatives" with '
                    '"International standardization initiatives"'
                ),
            ),
        ),
    )
    msg7 = container.repos.messages.put(
        Message(
            id="m7",
            session_id="ses_edit",
            role=Role.user,
            content=(
                'In the first section, replace "Global monitoring initiatives" with '
                '"International standardization initiatives"'
            ),
        )
    )
    reply7 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg7
    )
    assert reply7.version.version == 7
    doc7 = Document(io.BytesIO(container.blobs.read_bytes(reply7.version.blob_path)))
    assert any("International standardization initiatives" in p.text for p in doc7.paragraphs)

    # 8. Document restyle (version 8)
    fake_llm._queues["interpreter"].clear()
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.modify_artifact,
            confidence=0.95,
            objective="Microplastics report",
            change=ChangeRequest(
                operation=ChangeOperation.restyle,
                scope=ChangeScope(level=ScopeLevel.style, targets=[]),
                instruction="Change document headings font to Georgia",
            ),
        ),
    )
    fake_llm.add("style_planner", StylePlan(changes=[StyleChange(property="heading_font", value="Georgia")]))
    msg8 = container.repos.messages.put(
        Message(
            id="m8",
            session_id="ses_edit",
            role=Role.user,
            content="Change document headings font to Georgia",
        )
    )
    reply8 = await GenerationService(container).handle_message(
        container.repos.sessions.require(session.id), msg8
    )
    assert reply8.version.version == 8
    assert reply8.version.ir.style.heading_font == "Georgia"


def test_docx_validator_detects_placeholder_in_table(tmp_path: Path):
    plan = _plan()
    doc_path = tmp_path / "with_placeholder.docx"
    result = render_document(plan, generated_tokens("x"), doc_path)

    doc = Document(str(doc_path))
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].text = "TODO: add the missing result"
    doc.save(str(doc_path))

    run = CheckRun()
    check_docx_file(run, doc_path, result, None, plan, None)
    findings = [f for f in run.findings if f.check_id == "content.placeholder_text"]
    assert len(findings) == 1
    assert findings[0].severity == Severity.critical
    assert findings[0].dimension == Dimension.content
    assert "TODO: add the missing result" in findings[0].message


def _create_sample_logo_png() -> bytes:
    img = Image.new("RGB", (120, 60), color=(20, 80, 180))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_docx_template_extraction_header_footer_and_logo(tmp_path: Path):
    tpl = Document()
    tpl.styles["Heading 1"].font.name = "Georgia"
    section = tpl.sections[0]
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    # Header with logo + company branding text
    hp = section.header.paragraphs[0]
    hp.text = "ACME Global Research | Internal Only"
    run = hp.add_run()
    run.add_picture(io.BytesIO(_create_sample_logo_png()), width=Inches(1.2))

    # Footer with confidentiality and page number
    fp = section.footer.paragraphs[0]
    fp.text = "Confidential - Proprietary"
    from docforge.render.docx.renderer import _add_field

    fp.add_run(" - Page ")
    _add_field(fp, "PAGE", "1")

    tpl.add_paragraph("Sample body text that should be removed")
    template_path = tmp_path / "branded_template.docx"
    tpl.save(str(template_path))

    manifest = analyze_docx(template_path)
    assert manifest.format.value == "docx"
    assert manifest.docx is not None
    info = manifest.docx
    assert len(info.sections) == 1
    sec_info = info.sections[0]
    assert "ACME Global Research" in sec_info.header_text
    assert "Confidential" in sec_info.footer_text
    assert sec_info.has_page_number is True
    assert sec_info.margins["left"] == int(Inches(1.25))

    # Verify logo extraction
    assert len(info.image_assets) >= 1
    logo_asset = next((a for a in info.image_assets if a.is_logo), None)
    assert logo_asset is not None
    assert logo_asset.location == "header"
    assert logo_asset.width_emu is not None and logo_asset.width_emu > 0


def test_docx_template_reuse_generates_document_with_branding(tmp_path: Path):
    tpl = Document()
    tpl.styles["Heading 1"].font.name = "Georgia"
    section = tpl.sections[0]
    section.header.paragraphs[0].text = "ACME Corp Header"
    run = section.header.paragraphs[0].add_run()
    run.add_picture(io.BytesIO(_create_sample_logo_png()), width=Inches(1.0))
    section.footer.paragraphs[0].text = "ACME Corp Footer Confidential"

    tpl.add_paragraph("OLD SAMPLE CONTENT TO REMOVE")
    template_path = tmp_path / "branded_tpl2.docx"
    tpl.save(str(template_path))

    manifest = analyze_docx(template_path)
    out_path = tmp_path / "branded_output.docx"
    render_document(
        _plan(["c1", "c2"]),
        generated_tokens("test"),
        out_path,
        dossier=_dossier(),
        template=template_path,
        template_manifest=manifest,
    )

    doc = Document(str(out_path))
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "OLD SAMPLE CONTENT" not in body_text
    assert "What the evidence shows" in body_text

    # Verify header and footer preservation
    assert doc.sections[0].header.paragraphs[0].text.startswith("ACME Corp Header")
    assert "<w:drawing" in doc.sections[0].header._element.xml
    assert "ACME Corp Footer Confidential" in doc.sections[0].footer.paragraphs[0].text


def test_docx_template_multi_section_and_different_first_page(tmp_path: Path):
    tpl = Document()
    sec1 = tpl.sections[0]
    sec1.different_first_page_header_footer = True
    sec1.first_page_header.paragraphs[0].text = "First Page Header Only"
    sec1.header.paragraphs[0].text = "Default Section 1 Header"
    sec1.footer.paragraphs[0].text = "Section 1 Footer"

    # Add second section
    sec2 = tpl.add_section()
    sec2.header.paragraphs[0].text = "Section 2 Distinct Header"
    sec2.footer.paragraphs[0].text = "Section 2 Distinct Footer"

    template_path = tmp_path / "multi_sec.docx"
    tpl.save(str(template_path))

    manifest = analyze_docx(template_path)
    assert len(manifest.docx.sections) == 2
    assert manifest.docx.sections[0].different_first_page is True
    assert manifest.docx.sections[0].first_page_header.text == "First Page Header Only"
    assert manifest.docx.sections[1].default_header.text == "Section 2 Distinct Header"

    out_path = tmp_path / "multi_sec_out.docx"
    render_document(
        _plan(),
        generated_tokens("test"),
        out_path,
        template=template_path,
        template_manifest=manifest,
    )
    doc = Document(str(out_path))
    assert len(doc.sections) == 2
    assert "Right-click and choose Update Field" not in "\n".join(p.text for p in doc.paragraphs)
    assert doc.sections[1].header.paragraphs[0].text == "Section 2 Distinct Header"
    assert doc.sections[1].footer.paragraphs[0].text == "Section 2 Distinct Footer"


def test_docx_template_cover_logo_extracted_and_reused(tmp_path: Path):
    tpl = Document()
    p = tpl.add_paragraph()
    run = p.add_run()
    run.add_picture(io.BytesIO(_create_sample_logo_png()), width=Inches(2.5))
    tpl.add_paragraph("Sample cover title")
    tpl.add_paragraph("Sample body")

    template_path = tmp_path / "cover_logo_tpl.docx"
    tpl.save(str(template_path))

    manifest = analyze_docx(template_path)
    assert any(a.is_logo for a in manifest.docx.image_assets)

    out_path = tmp_path / "cover_logo_out.docx"
    render_document(
        _plan(),
        generated_tokens("test"),
        out_path,
        template=template_path,
        template_manifest=manifest,
    )
    doc = Document(str(out_path))
    # Check that a drawing was added to the title page area
    assert any("<w:drawing" in p._element.xml for p in doc.paragraphs[:5])


def test_docx_template_no_logo_renders_cleanly(tmp_path: Path):
    tpl = Document()
    tpl.styles["Normal"].font.name = "Arial"
    tpl.sections[0].header.paragraphs[0].text = "Header Without Logo"
    tpl.add_paragraph("Body paragraph")
    template_path = tmp_path / "no_logo_tpl.docx"
    tpl.save(str(template_path))

    manifest = analyze_docx(template_path)
    assert not any(a.is_logo for a in manifest.docx.image_assets)

    out_path = tmp_path / "no_logo_out.docx"
    render_document(
        _plan(),
        generated_tokens("test"),
        out_path,
        template=template_path,
        template_manifest=manifest,
    )
    doc = Document(str(out_path))
    assert doc.sections[0].header.paragraphs[0].text == "Header Without Logo"
    assert "PAGE" in doc.sections[0].footer._element.xml


async def test_generation_service_generates_docx_with_template(container, fake_llm, tmp_path: Path):
    tpl = Document()
    tpl.styles["Normal"].font.name = "Arial"
    tpl.sections[0].header.paragraphs[0].text = "Acme Annual Report"
    run = tpl.sections[0].header.paragraphs[0].add_run()
    run.add_picture(io.BytesIO(_create_sample_logo_png()), width=Inches(1.0))
    template_path = tmp_path / "acme_template.docx"
    tpl.save(str(template_path))

    blob = container.blobs.put_file(template_path)
    session = container.repos.sessions.put(Session(id="ses_tpl_doc", title="Acme Session", owner="demo"))
    container.repos.files.put(
        UploadedFile(
            id="f_tpl",
            session_id=session.id,
            filename="acme_template.docx",
            role=FileRole.template,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size=len(blob.path),
            sha256="fake_sha",
            blob_path=blob.path,
        )
    )

    fake_llm.add(
        "interpreter",
        Interpretation(
            kind=RequestKind.new_task,
            confidence=0.95,
            objective="Generate an Acme report",
            output_spec=OutputSpec(format=OutputFormat.docx),
        ),
    )
    fake_llm.add("doc_outliner", outline())
    for w in written():
        fake_llm.add("section_writer", w)

    msg = container.repos.messages.put(
        Message(id="m_tpl", session_id=session.id, role=Role.user, content="Generate an Acme report")
    )
    reply = await GenerationService(container).handle_message(session, msg)
    assert reply.version is not None
    assert reply.version.version == 1
    doc = Document(io.BytesIO(container.blobs.read_bytes(reply.version.blob_path)))
    assert "Acme Annual Report" in doc.sections[0].header.paragraphs[0].text
    assert "<w:drawing" in doc.sections[0].header._element.xml
