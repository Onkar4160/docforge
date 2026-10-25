from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from docforge.design.system import contrast_ratio, generated_tokens, palette_names, readable_on
from docforge.layout.text import EMU_PER_INCH, Para, fit_paragraphs, fits, measure
from docforge.models.intent import LengthSpec, OutputFormat, OutputSpec
from docforge.models.ir import BulletItem, BulletsBlock, DeckPlan, SlidePlan
from docforge.models.quality import FindingStatus, Verdict
from docforge.render.pptx import render_deck
from docforge.render.pptx.composer import effective_archetype
from docforge.render.pptx.preview import render_previews
from docforge.validate.common import CheckRun
from docforge.validate.deck import build_deck, validate_deck_file
from docforge.validate.plan_checks import check_deck_plan
from docforge.validate.pptx_checks import check_pptx_file
from docforge.validate.repair import repair_deck, split_slide
from tests.fixtures.decks import sample_deck

INCH = EMU_PER_INCH


def _tokens(deck: DeckPlan):
    return generated_tokens(deck.title, hint=deck.design_hint)


def _bullets_slide(sid: str, n: int, words: int = 8) -> SlidePlan:
    items = [BulletItem(text=" ".join([f"point{i}"] + ["detail"] * words)) for i in range(n)]
    return SlidePlan(
        id=sid,
        purpose="p",
        archetype="bullets",
        headline=f"Headline {sid}",
        blocks=[BulletsBlock(items=items)],
    )


# ---------------------------------------------------------------- text measurement (ADR-010)


def test_measure_grows_with_text_and_narrower_boxes():
    short = measure([Para("one line", 18)], 6 * INCH)
    long = measure([Para("word " * 80, 18)], 6 * INCH)
    narrow = measure([Para("word " * 80, 18)], 3 * INCH)
    assert short.lines == 1 and long.lines > 1
    assert narrow.height_emu > long.height_emu > short.height_emu


def test_fit_paragraphs_shrinks_until_fit_and_respects_floor():
    paras = [Para("Heading", 28, bold=True), Para("word " * 60, 20)]
    fitted, ok = fit_paragraphs(paras, 6 * INCH, 2 * INCH, min_pt=10)
    assert ok and fitted[1].size_pt < 20 and fitted[0].size_pt > fitted[1].size_pt
    assert fits(fitted, 6 * INCH, 2 * INCH)
    huge = [Para("word " * 800, 20)]
    fitted, ok = fit_paragraphs(huge, 3 * INCH, 1 * INCH, min_pt=14)
    assert not ok and fitted[0].size_pt == 14


def test_word_wider_than_box_is_flagged():
    assert measure([Para("Supercalifragilisticexpialidocious", 40)], int(0.5 * INCH)).widest_word_overflow


# ---------------------------------------------------------------- design tokens (FR-PPTX-3)


def test_palette_choice_by_hint_keywords_and_stable_hash():
    assert generated_tokens("x", hint="teal_trust").palette.primary == "028090"
    health = generated_tokens("Clinical trial outcomes for hospital boards")
    assert health.palette.primary == "028090"
    a, b = generated_tokens("Quarterly widget review"), generated_tokens("Quarterly widget review")
    assert a == b and a.source == "generated"
    assert "midnight_executive" in palette_names()


def test_contrast_helpers():
    assert contrast_ratio("FFFFFF", "000000") == pytest.approx(21, rel=0.01)
    assert readable_on("1E2761") == "FFFFFF"
    assert readable_on("F5F5F5") == "1F2328"


# ---------------------------------------------------------------- rendering (FR-PPTX-1/4/5/6)


def test_sample_deck_renders_every_archetype_cleanly(tmp_path: Path):
    deck = sample_deck()
    build = build_deck(
        deck, _tokens(deck), tmp_path, spec=None, dossier=None, max_iterations=3, libreoffice=None
    )
    report = build.report
    unresolved = [f for f in report.findings if f.status != FindingStatus.corrected]
    assert unresolved == [], [f.message for f in unresolved]
    assert report.verdict == Verdict.pass_with_warnings  # only skipped tier-4/5 checks
    assert {c.check_id for c in report.checks if c.status == "skipped"} == {
        "common.convert",
        "pptx.visual_clipping",
        "content.review",
    }

    prs = Presentation(str(build.path))
    assert len(prs.slides) == len(deck.slides)
    assert prs.slide_width == int(13.333 * INCH)
    names = [{s.name for s in slide.shapes} for slide in prs.slides]
    assert "content:chart" in names[3] and "content:table" in names[8]
    texts = " ".join(sh.text_frame.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame)
    assert "•" not in texts  # bullets are paragraph properties, not characters
    assert prs.slides[3].notes_slide is not None


def test_render_is_structurally_deterministic(tmp_path: Path):
    deck = sample_deck()

    def dump(path: Path) -> list:
        prs = Presentation(str(path))
        return [
            [
                (s.name, s.left, s.top, s.width, s.height, s.text_frame.text if s.has_text_frame else None)
                for s in slide.shapes
            ]
            for slide in prs.slides
        ]

    a = render_deck(deck, _tokens(deck), tmp_path / "a.pptx").path
    b = render_deck(deck, _tokens(deck), tmp_path / "b.pptx").path
    assert dump(a) == dump(b)


def test_effective_archetype_falls_back_when_block_missing():
    slide = SlidePlan(
        id="s1",
        purpose="p",
        archetype="chart",
        headline="h",
        blocks=[BulletsBlock(items=[BulletItem(text="a")])],
    )
    assert effective_archetype(slide).value == "bullets"


def test_previews_render(tmp_path: Path):
    deck = sample_deck()
    path = render_deck(deck, _tokens(deck), tmp_path / "d.pptx").path
    images = render_previews(path, tmp_path / "png", width_px=640)
    assert len(images) == len(deck.slides) and all(p.stat().st_size > 1000 for p in images)


# ---------------------------------------------------------------- validator catches seeded defects (FR-VAL-1)


def _defect_deck(tmp_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5))
    box.name = "content:body"
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "far too much text " * 30
    run.font.size = Pt(18)
    off = slide.shapes.add_textbox(Inches(12.5), Inches(6.9), Inches(2), Inches(1))
    off.name = "content:offslide"
    off.text_frame.text = "Lorem ipsum dolor"
    tiny = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    tiny.name = "content:tiny"
    r = tiny.text_frame.paragraphs[0].add_run()
    r.text = "small"
    r.font.size = Pt(9)
    path = tmp_path / "defects.pptx"
    prs.save(str(path))
    return path


def test_validator_detects_seeded_defects(tmp_path: Path):
    run = CheckRun()
    check_pptx_file(run, _defect_deck(tmp_path), ["s01"], libreoffice=None)
    found = {f.check_id for f in run.findings}
    assert {
        "pptx.text_overflow",
        "pptx.out_of_bounds",
        "content.placeholder_text",
        "pptx.font_floor",
    } <= found
    assert all(f.unit_id == "s01" for f in run.findings)


def test_validator_reports_unopenable_file(tmp_path: Path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a zip")
    run = CheckRun()
    check_pptx_file(run, bad, [], libreoffice=None)
    assert [f.check_id for f in run.findings] == ["common.reopen"]


def test_plan_checks_length_duplicates_empty_and_citations():
    deck = DeckPlan(
        title="t",
        slides=[
            _bullets_slide("s01", 3),
            SlidePlan(id="s02", purpose="p", archetype="bullets", headline="Headline s01"),
            SlidePlan(
                id="s03",
                purpose="p",
                archetype="bullets",
                headline="Other",
                claim_ids=["clm_x"],
                blocks=[BulletsBlock(items=[BulletItem(text="a")])],
            ),
        ],
    )
    spec = OutputSpec(format=OutputFormat.pptx, length=LengthSpec(unit="slides", target=5, explicit=True))
    run = CheckRun()
    check_deck_plan(run, deck, spec, dossier=None)
    found = {f.check_id for f in run.findings}
    assert {
        "common.length",
        "content.duplicates",
        "content.empty_units",
        "evidence.citations_resolve",
    } <= found


# ---------------------------------------------------------------- repair loop (FR-VAL-2)


def test_split_slide_halves_items():
    parts = split_slide(_bullets_slide("s05", 7))
    assert [len(p.blocks[0].items) for p in parts] == [4, 3]
    assert parts[1].id == "s05-2" and parts[1].headline.endswith("(continued)")


def test_overflowing_slide_is_split_and_marked_corrected(tmp_path: Path):
    deck = DeckPlan(title="t", slides=[_bullets_slide("s01", 12, words=30)])
    tokens = generated_tokens("t")
    first = render_deck(deck, tokens, tmp_path / "first.pptx").path
    assert any(
        f.check_id == "pptx.text_overflow" for f in validate_deck_file(first, deck, None, None, None).findings
    )

    build = build_deck(deck, tokens, tmp_path, spec=None, dossier=None, max_iterations=3, libreoffice=None)
    assert len(build.deck.slides) > 1
    assert any(f.status == FindingStatus.corrected and f.repair_action for f in build.report.findings)
    assert build.report.iterations >= 1


def test_fixed_length_blocks_splitting(tmp_path: Path):
    deck = DeckPlan(title="t", slides=[_bullets_slide("s01", 12, words=30)])
    run = CheckRun()
    run.finding("pptx.text_overflow", "layout", "critical", "overflow", unit_id="s01")
    same, actions = repair_deck(deck, run.findings, allow_growth=False)
    assert same == deck and actions == []
    spec = OutputSpec(format=OutputFormat.pptx, length=LengthSpec(unit="slides", target=1, explicit=True))
    build = build_deck(
        deck, generated_tokens("t"), tmp_path, spec=spec, dossier=None, max_iterations=3, libreoffice=None
    )
    assert len(build.deck.slides) == 1 and build.report.verdict == Verdict.failed
