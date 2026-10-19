from __future__ import annotations

import io
import logging
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from docforge.core.config import Settings
from docforge.core.errors import LLMRefused, UnsupportedFileError
from docforge.core.files import sniff
from docforge.core.ids import new_id
from docforge.core.logging import JsonFormatter, _ContextFilter, get_logger, log_context

# ---------------------------------------------------------------- config


def test_settings_read_project_env_names(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    for key, value in {
        "GEMINI_API_KEY": "g-key",
        "TAVILY_API_KEY": "t-key",
        "REQUIRE_AUTHENTICATED_WEB_SEARCH": "false",
        "JWT_EXPIRE_MINUTES": "60",
        "UPLOAD_MAX_MB": "10",
        "LLM_PRIMARY_MODEL": "gemini-2.5-pro",
        "LLM_FALLBACK_MODEL": "openai_compat:llama3.1",
        "LLM_MAX_RETRIES": "5",
        "PPTX_MASTERS_ENABLED": "1",
        "DATABASE_URL": "sqlite:///./x/app.db",
    }.items():
        monkeypatch.setenv(key, value)
    s = Settings(_env_file=None)
    assert s.gemini_api_key == "g-key" and s.tavily_api_key == "t-key"
    assert s.require_authenticated_web_search is False
    assert s.jwt_expire_minutes == 60 and s.max_upload_bytes == 10 * 1024 * 1024
    assert s.llm_primary_model == "gemini-2.5-pro" and s.llm_fallback_model == "openai_compat:llama3.1"
    assert s.llm_max_retries == 5 and s.pptx_masters_enabled is True


def test_placeholder_keys_count_as_unset(tmp_path: Path):
    s = Settings(
        _env_file=None, gemini_api_key="your-gemini-api-key-here", pinecone_api_key="", tavily_api_key=" "
    )
    assert s.gemini_api_key is None and s.pinecone_api_key is None and s.tavily_api_key is None
    assert s.resolved_vector_store == "local"
    assert Settings(_env_file=None, pinecone_api_key="pc-1").resolved_vector_store == "pinecone"


def test_settings_split_domain_lists():
    s = Settings(_env_file=None, research_allowed_domains="a.org, b.com ,")
    assert s.research_allowed_domains == ["a.org", "b.com"]


def test_settings_per_agent_model_override():
    s = Settings(_env_file=None, llm_agent_model={"planner": "gemini-2.5-pro"})
    assert s.model_for("planner") == "gemini-2.5-pro"
    assert s.model_for("interpreter") == "gemini-3.8-flash"


def test_default_template_only_when_file_exists(tmp_path: Path):
    real = tmp_path / "t.pptx"
    real.write_bytes(b"x")
    s = Settings(
        _env_file=None, default_pptx_template_path=real, default_docx_template_path=tmp_path / "no.docx"
    )
    assert s.default_template("pptx") == real
    assert s.default_template("docx") is None
    assert s.default_template("xlsx") is None


def test_ensure_dirs_creates_sqlite_parent(tmp_path: Path):
    s = Settings(
        _env_file=None,
        storage_dir=tmp_path / "st",
        database_url=f"sqlite:///{(tmp_path / 'db' / 'a.db').as_posix()}",
    )
    s.ensure_dirs()
    assert (tmp_path / "db").is_dir() and (tmp_path / "st" / "blobs").is_dir()


def test_missing_libreoffice_path_returns_none(tmp_path: Path):
    s = Settings(_env_file=None, libreoffice_path=tmp_path / "nope" / "soffice")
    assert s.find_libreoffice() is None


# ---------------------------------------------------------------- ids


def test_ids_prefixed_unique_and_time_sortable():
    ids = [new_id("session") for _ in range(200)]
    assert all(i.startswith("ses_") and len(i) == 30 for i in ids)
    assert len(set(ids)) == 200
    first = new_id("task")
    time.sleep(0.002)
    later = new_id("task")
    assert first[:14] <= later[:14]


# ---------------------------------------------------------------- errors / logging


def test_refusal_error_carries_category():
    err = LLMRefused("declined", category="cyber")
    assert err.to_dict()["details"]["category"] == "cyber"
    assert err.http_status == 502


def test_log_context_in_json_output():
    logger = get_logger("test")
    record = logger.makeRecord(logger.name, logging.INFO, __file__, 1, "hello %s", ("x",), None)
    with log_context(job_id="job_1"):
        _ContextFilter().filter(record)
    out = JsonFormatter().format(record)
    assert '"job_id": "job_1"' in out and '"msg": "hello x"' in out


# ---------------------------------------------------------------- file sniffing (NFR-6)


def _pptx_bytes() -> bytes:
    buf = io.BytesIO()
    Presentation().save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    buf = io.BytesIO()
    Document().save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    buf = io.BytesIO()
    Workbook().save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(
    ("name", "factory", "family"),
    [
        ("deck.pptx", _pptx_bytes, "pptx"),
        ("doc.docx", _docx_bytes, "docx"),
        ("book.xlsx", _xlsx_bytes, "xlsx"),
    ],
)
def test_sniff_ooxml(name, factory, family):
    kind = sniff(name, factory())
    assert kind.family == family and kind.is_template_capable


def test_sniff_rejects_mismatched_ooxml():
    with pytest.raises(UnsupportedFileError, match="does not contain"):
        sniff("deck.pptx", _docx_bytes())


def test_sniff_rejects_fake_pdf_and_unknown_extension():
    with pytest.raises(UnsupportedFileError):
        sniff("x.pdf", b"not a pdf")
    with pytest.raises(UnsupportedFileError, match="unsupported"):
        sniff("x.exe", b"MZ")


def test_sniff_rejects_macro_package():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("ppt/presentation.xml", "<p/>")
        zf.writestr("ppt/vbaProject.bin", b"\0")
    with pytest.raises(UnsupportedFileError, match="macro"):
        sniff("deck.pptx", buf.getvalue())


def test_sniff_rejects_zip_bomb_ratio():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/huge.xml", b"0" * 20_000_000)
    with pytest.raises(UnsupportedFileError, match="safe limits"):
        sniff("book.xlsx", buf.getvalue())


def test_sniff_text_requires_utf8():
    assert sniff("data.csv", b"a,b\n1,2\n").family == "text"
    with pytest.raises(UnsupportedFileError):
        sniff("data.csv", b"\xff\xfe\x00bad")
