"""User-shaped generation and precise-edit regression tests."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation

from docforge.agents.doc_writer import assemble
from docforge.agents.editor import exact_section_replacement
from docforge.design.system import generated_tokens
from docforge.orchestrator.patch import apply_slide_edit
from docforge.render.docx import render_document
from docforge.validate.deck import build_deck
from tests.fixtures.decks import sample_deck
from tests.fixtures.documents import outline, written


def sample_document():
    return assemble(outline(), written())


def test_real_world_document_edit_replaces_one_sentence_and_keeps_the_rest(tmp_path: Path):
    original = sample_document()
    target = original.sections[0]
    instruction = 'Replace the caption "Detection methods" with "Detection methods used in the review".'

    edited = exact_section_replacement(target, instruction)

    assert edited is not None
    assert edited.blocks[2].caption == "Detection methods used in the review"
    assert edited.blocks[0] == target.blocks[0]
    assert edited.blocks[1] == target.blocks[1]
    assert edited.blocks[3:] == target.blocks[3:]
    assert original.sections[1] == sample_document().sections[1]

    output = tmp_path / "edited-report.docx"
    render_document(
        original.model_copy(update={"sections": [edited, *original.sections[1:]]}),
        generated_tokens(original.title),
        output,
    )
    paragraphs = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
    assert "Detection methods used in the review" in paragraphs
    assert "Options for the authority" in paragraphs


def test_real_world_presentation_edit_changes_one_slide_only(tmp_path: Path):
    original = sample_deck()
    edited_slide = original.slides[3].model_copy(
        update={"headline": "Public charge points will expand by 2028"}
    )
    edited, patch = apply_slide_edit(original, edited_slide, instruction="update the chart headline")

    assert patch.affected_unit_ids == [original.slides[3].id]
    assert edited.slides[3].headline == "Public charge points will expand by 2028"
    assert edited.slides[:3] == original.slides[:3]
    assert edited.slides[4:] == original.slides[4:]

    build = build_deck(
        edited,
        generated_tokens(edited.title, hint=edited.design_hint),
        tmp_path,
        spec=None,
        dossier=None,
        max_iterations=3,
        libreoffice=None,
    )
    presentation = Presentation(str(build.path))
    slide_text = " ".join(
        shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Public charge points will expand by 2028" in slide_text
    assert len(presentation.slides) == len(original.slides)