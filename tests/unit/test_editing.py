"""Unit tests for targeted editing, target resolution, and Scope Guard (FR-EDIT-1..4)."""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

from docforge.agents.doc_writer import assemble
from docforge.agents.editor import (
    SectionEditorAgent,
    SectionEditorInput,
    SectionInsertionAgent,
    SectionInsertionInput,
    SlideEditorAgent,
    SlideEditorInput,
    SlideInsertionAgent,
    SlideInsertionInput,
)
from docforge.agents.style_planner import StylePlannerAgent, StylePlannerInput
from docforge.llm import FakeLLMClient
from docforge.models.design import DesignTokens
from docforge.models.intent import (
    SlidePlacement,
    SlidePlacementMode,
    StyleChange,
    StylePlan,
    TargetKind,
    TargetRef,
)
from docforge.models.ir import (
    DocumentPlan,
    ParagraphBlock,
    SectionDraft,
    SectionInsertionPlan,
    SectionPlan,
    SlideDraft,
    SlideInsertionPlan,
    SlidePlan,
    StyleSpec,
)
from docforge.orchestrator.patch import (
    ScopeViolationError,
    apply_remove_unit,
    apply_restyle,
    apply_section_block_edit,
    apply_section_edit,
    apply_section_insert,
    apply_section_move,
    apply_sheet_edit,
    apply_slide_edit,
    apply_slide_insert,
    apply_slide_move,
)
from docforge.orchestrator.targets import (
    resolve_section_position,
    resolve_slide_position,
    resolve_targets,
)
from docforge.render.docx import render_document
from tests.fixtures.decks import sample_deck
from tests.fixtures.documents import outline, written
from tests.fixtures.workbooks import saas_model

pytestmark = pytest.mark.anyio


def sample_document() -> DocumentPlan:
    return assemble(outline(), written())


def test_target_resolver_by_index_and_ordinal():
    deck = sample_deck()  # 13 slides
    res1 = resolve_targets(deck, None, "change slide 3 to focus on battery tech")
    assert not res1.clarification_needed
    assert len(res1.targets) == 1
    assert res1.targets[0].unit_id == deck.slides[2].id

    res2 = resolve_targets(deck, None, "update the fourth slide")
    assert not res2.clarification_needed
    assert res2.targets[0].unit_id == deck.slides[3].id

    res3 = resolve_targets(deck, None, "change the last slide")
    assert res3.targets[0].unit_id == deck.slides[-1].id

    res4 = resolve_targets(deck, None, "update slide #1")
    assert res4.targets[0].unit_id == deck.slides[0].id


def test_target_resolver_by_headline_or_keyword():
    deck = sample_deck()
    # Find a slide with known headline word
    target_slide = deck.slides[4]
    res = resolve_targets(deck, None, f"on the '{target_slide.headline}', make bullets shorter")
    assert res.targets and res.targets[0].unit_id == target_slide.id


def test_target_resolver_document_sections():
    doc = sample_document()
    all_secs = doc.all_sections()
    res = resolve_targets(doc, None, "change section 2 to mention Europe")
    assert res.targets and res.targets[0].unit_id == all_secs[1].id


def test_target_resolver_workbook_sheets():
    wb = saas_model()
    res = resolve_targets(wb, None, "update the Assumptions sheet")
    assert res.targets and res.targets[0].unit_id == "a"


def test_target_resolver_ambiguous_prompts_clarification():
    deck = sample_deck()
    res = resolve_targets(deck, None, "make it better and change the points")
    assert res.clarification_needed
    assert "Which slide" in res.clarification_question


def test_apply_slide_edit_preserves_untouched_slides():
    deck = sample_deck()
    original_slides = [s.model_copy(deep=True) for s in deck.slides]

    # Modify slide 2
    target_id = deck.slides[1].id
    edited = deck.slides[1].model_copy(
        update={"headline": "Updated Battery Technology", "key_message": "New 2026 specs"}
    )
    new_deck, patch = apply_slide_edit(deck, edited, instruction="update headline", base_version=1)

    assert patch.affected_unit_ids == [target_id]
    assert new_deck.slides[1].headline == "Updated Battery Technology"

    # Verify all other 12 slides are 100% identical
    for i, s in enumerate(new_deck.slides):
        if i != 1:
            assert s == original_slides[i]


def test_apply_slide_edit_rejects_missing_target():
    deck = sample_deck()
    fake_slide = SlidePlan(id="non_existent", purpose="p", archetype="bullets", headline="h")
    with pytest.raises(ScopeViolationError):
        apply_slide_edit(deck, fake_slide)


def test_insert_and_move_keep_existing_slide_ids_and_order():
    deck = sample_deck()
    original_ids = [slide.id for slide in deck.slides]
    inserted = SlidePlan(
        id="slide_new",
        purpose="bridge",
        archetype="bullets",
        headline="A new bridge slide",
    )

    with_insert, insert_patch = apply_slide_insert(deck, inserted, 2, instruction="insert before slide 3")
    assert [slide.id for slide in with_insert.slides] == [*original_ids[:2], "slide_new", *original_ids[2:]]
    assert insert_patch.ops[0].path == "/slides/2"
    assert [slide.id for slide in deck.slides] == original_ids

    moved, move_patch = apply_slide_move(with_insert, "s01", len(with_insert.slides) - 1)
    assert moved.slides[-1].id == "s01"
    assert set(slide.id for slide in moved.slides) == set(slide.id for slide in with_insert.slides)
    assert moved.slides[0].id == "s02"
    assert move_patch.ops[0].op == "move"


def test_slide_position_requires_adjacent_between_anchors():
    deck = sample_deck()
    result = resolve_slide_position(
        deck,
        SlidePlacement(
            mode=SlidePlacementMode.between,
            before=TargetRef(kind=TargetKind.slide, locator="slide 2"),
            after=TargetRef(kind=TargetKind.slide, locator="slide 4"),
        ),
    )
    assert result.clarification_needed
    assert "not adjacent" in result.clarification_question
    assert len(result.options) == 2


def test_slide_position_resolves_start_end_and_adjacent_anchors():
    deck = sample_deck()
    start = resolve_slide_position(deck, SlidePlacement(mode=SlidePlacementMode.start))
    end = resolve_slide_position(deck, SlidePlacement(mode=SlidePlacementMode.end))
    middle = resolve_slide_position(
        deck,
        SlidePlacement(
            mode=SlidePlacementMode.between,
            before=TargetRef(kind=TargetKind.slide, locator="slide 2"),
            after=TargetRef(kind=TargetKind.slide, locator="slide 3"),
        ),
    )
    assert (start.index, end.index, middle.index) == (0, len(deck.slides), 2)


def test_apply_restyle_can_target_slides_and_rejects_unknown_properties():
    deck = sample_deck()
    target_id = deck.slides[2].id
    styled, patch = apply_restyle(
        deck,
        [StyleChange(property="heading_font", value="Georgia")],
        slide_ids=[target_id],
    )
    assert styled.slides[2].style.heading_font == "Georgia"
    assert styled.slides[:2] == deck.slides[:2]
    assert styled.slides[3:] == deck.slides[3:]
    assert patch.affected_unit_ids == [target_id]
    with pytest.raises(ScopeViolationError, match="Unsupported style property"):
        apply_restyle(deck, [StyleChange(property="arbitrary_field", value="x")])


def test_apply_section_edit_preserves_other_sections():
    doc = sample_document()
    target_sec = doc.sections[0]
    edited = target_sec.model_copy(update={"heading": "New Executive Overview"})
    new_doc, patch = apply_section_edit(doc, edited, instruction="rename heading", base_version=1)

    assert patch.affected_unit_ids == [target_sec.id]
    assert new_doc.sections[0].heading == "New Executive Overview"
    assert new_doc.sections[1:] == doc.sections[1:]


def test_apply_section_block_edit_preserves_every_other_document_chunk():
    doc = sample_document()
    target = doc.sections[0]
    edited = target.blocks[0].model_copy(update={"text": "Only this paragraph changed."})
    new_doc, patch = apply_section_block_edit(
        doc, target.id, 0, edited, instruction="rewrite the first paragraph", base_version=1
    )

    assert patch.ops[0].path == f"/sections/{target.id}/blocks/0"
    assert new_doc.sections[0].blocks[0] == edited
    assert new_doc.sections[0].blocks[1:] == target.blocks[1:]
    assert new_doc.sections[1:] == doc.sections[1:]


def test_apply_sheet_edit_preserves_other_sheets():
    wb = saas_model()
    target_sheet = wb.sheets[0]
    edited = target_sheet.model_copy(update={"description": "Updated assumptions"})
    new_wb, patch = apply_sheet_edit(wb, edited, instruction="update desc", base_version=1)

    assert patch.affected_unit_ids == [target_sheet.id]
    assert new_wb.sheets[0].description == "Updated assumptions"
    assert new_wb.sheets[1:] == wb.sheets[1:]


def test_apply_remove_unit():
    deck = sample_deck()
    remove_id = deck.slides[2].id
    new_deck, patch = apply_remove_unit(deck, remove_id, instruction="delete slide 3", base_version=1)

    assert len(new_deck.slides) == len(deck.slides) - 1
    assert remove_id not in [s.id for s in new_deck.slides]


def test_apply_restyle():
    deck = sample_deck()
    changes = [StyleChange(property="heading_font", value="Georgia")]
    new_deck, patch = apply_restyle(deck, changes, instruction="change font to Georgia", base_version=1)

    assert new_deck.style.heading_font == "Georgia"
    assert new_deck.slides == deck.slides  # content completely untouched


async def test_slide_editor_agent():
    deck = sample_deck()
    slide = deck.slides[1]
    updated_slide = slide.model_copy(update={"headline": "Edited by Agent"})
    fake = FakeLLMClient({"slide_editor": updated_slide})

    agent = SlideEditorAgent(fake)
    result = await agent.run(
        SlideEditorInput(
            slide=slide,
            instruction="change headline",
            outline=[(s.id, s.headline) for s in deck.slides],
        )
    )
    assert result.headline == "Edited by Agent"
    assert result.id == slide.id


async def test_slide_insertion_agent_receives_position_and_neighbor_context():
    deck = sample_deck()
    draft = SlideDraft(purpose="bridge", archetype="bullets", headline="Bridge the argument")
    fake = FakeLLMClient({"slide_inserter": SlideInsertionPlan(slides=[draft])})
    agent = SlideInsertionAgent(fake)

    result = await agent.run(
        SlideInsertionInput(
            deck=deck,
            instruction="Add a transition slide",
            insertion_index=2,
            template_guidance="Preserve the observed template typography.",
        )
    )

    prompt = fake.calls_for("slide_inserter")[0].messages[0].content
    assert result.slides[0].headline == "Bridge the argument"
    assert "Index: 2" in prompt
    assert deck.slides[1].headline in prompt and deck.slides[2].headline in prompt
    assert "Preserve the observed template typography" in prompt
    assert not hasattr(result.slides[0], "id")


async def test_style_planner_receives_scope_tokens_and_template_rules():
    deck = sample_deck()
    plan = StylePlan(changes=[StyleChange(property="heading_font", value="Georgia")])
    fake = FakeLLMClient({"style_planner": plan})
    agent = StylePlannerAgent(fake)

    result = await agent.run(
        StylePlannerInput(
            deck=deck,
            tokens=DesignTokens(),
            instruction="Use Georgia for headings",
            scope_description="slides s02, s03",
            target_slide_ids=["s02", "s03"],
        )
    )

    prompt = fake.calls_for("style_planner")[0].messages[0].content
    assert result.changes[0].value == "Georgia"
    assert "slides s02, s03" in prompt
    assert "Design tokens" in prompt


def test_resolve_section_position():
    doc = sample_document()
    assert len(doc.sections) >= 2
    start = resolve_section_position(doc, SlidePlacement(mode=SlidePlacementMode.start))
    end = resolve_section_position(doc, SlidePlacement(mode=SlidePlacementMode.end))
    after_first = resolve_section_position(
        doc,
        SlidePlacement(
            mode=SlidePlacementMode.after,
            anchor=TargetRef(kind=TargetKind.section, locator="section 1"),
        ),
    )
    before_second = resolve_section_position(
        doc,
        SlidePlacement(
            mode=SlidePlacementMode.before,
            anchor=TargetRef(kind=TargetKind.section, locator="section 2"),
        ),
    )
    between_adjacent = resolve_section_position(
        doc,
        SlidePlacement(
            mode=SlidePlacementMode.between,
            before=TargetRef(kind=TargetKind.section, locator="section 1"),
            after=TargetRef(kind=TargetKind.section, locator="section 2"),
        ),
    )
    assert (start.index, end.index, after_first.index, before_second.index, between_adjacent.index) == (
        0,
        len(doc.sections),
        1,
        1,
        1,
    )


def test_resolve_section_position_non_adjacent_requires_clarification():
    doc = sample_document()
    extra_section = SectionPlan(
        id="sec_extra",
        heading="Extra Intermediary Section",
        purpose="testing",
        level=1,
    )
    doc_three, _ = apply_section_insert(doc, extra_section, 1)
    assert len(doc_three.sections) == 3

    non_adjacent = resolve_section_position(
        doc_three,
        SlidePlacement(
            mode=SlidePlacementMode.between,
            before=TargetRef(kind=TargetKind.section, locator="section 1"),
            after=TargetRef(kind=TargetKind.section, locator="section 3"),
        ),
    )
    assert non_adjacent.clarification_needed
    assert "not adjacent" in non_adjacent.clarification_question
    assert len(non_adjacent.options) == 2


def test_apply_section_insert_and_move():
    doc = sample_document()
    original_ids = [s.id for s in doc.sections]
    new_sec = SectionPlan(
        id="sec_new_annex",
        heading="Annex: Risk Methodology",
        purpose="methodology",
        level=1,
        blocks=[ParagraphBlock(text="Methodology details.")],
    )

    # Insert at position 1
    with_insert, insert_patch = apply_section_insert(
        doc, new_sec, 1, instruction="add annex section after section 1"
    )
    assert [s.id for s in with_insert.sections] == [original_ids[0], "sec_new_annex", original_ids[1]]
    assert insert_patch.affected_unit_ids == ["sec_new_annex"]
    assert insert_patch.ops[0].path == "/sections/1"

    # Reorder / Move section 0 to end
    moved, move_patch = apply_section_move(with_insert, original_ids[0], len(with_insert.sections) - 1)
    assert moved.sections[-1].id == original_ids[0]
    assert moved.sections[0].id == "sec_new_annex"
    assert move_patch.ops[0].op == "move"

    # Out of bounds index raises ScopeViolationError
    with pytest.raises(ScopeViolationError):
        apply_section_move(doc, original_ids[0], 99)


def test_apply_remove_unit_document():
    doc = sample_document()
    top_id = doc.sections[0].id
    new_doc, patch = apply_remove_unit(doc, top_id, instruction="remove opening section", base_version=1)
    assert len(new_doc.sections) == len(doc.sections) - 1
    assert top_id not in [s.id for s in new_doc.sections]
    assert patch.affected_unit_ids == [top_id]

    # Remove subsection
    sec_with_subs = next(s for s in doc.sections if s.subsections)
    sub_id = sec_with_subs.subsections[0].id
    sub_count_before = len(sec_with_subs.subsections)
    doc_after_sub_remove, sub_patch = apply_remove_unit(
        doc, sub_id, instruction="remove first subsection", base_version=1
    )
    updated_parent = next(s for s in doc_after_sub_remove.sections if s.id == sec_with_subs.id)
    assert len(updated_parent.subsections) == sub_count_before - 1
    assert sub_id not in [sub.id for sub in updated_parent.subsections]
    assert sub_patch.affected_unit_ids == [sub_id]

    # Non-existent section raises ScopeViolationError
    with pytest.raises(ScopeViolationError):
        apply_remove_unit(doc, "sec_non_existent")


def test_apply_restyle_document_global_and_scoped():
    doc = sample_document()
    # Global document restyle
    changes = [
        StyleChange(property="heading_font", value="Georgia"),
        StyleChange(property="body_font", value="Consolas"),
        StyleChange(property="body_size_pt", value="12"),
        StyleChange(property="margins_in", value="1.25"),
        StyleChange(property="page_orientation", value="landscape"),
        StyleChange(property="primary", value="#003366"),
    ]
    styled_doc, patch = apply_restyle(doc, changes, instruction="apply executive theme")
    assert styled_doc.style.heading_font == "Georgia"
    assert styled_doc.style.body_font == "Consolas"
    assert styled_doc.style.body_size_pt == 12.0
    assert styled_doc.style.margins_in == 1.25
    assert styled_doc.style.page_orientation == "landscape"
    assert styled_doc.style.palette_overrides["primary"] == "003366"
    assert styled_doc.sections == doc.sections  # Content untouched

    # Section-scoped restyle
    target_sec = doc.sections[0]
    sec_changes = [StyleChange(property="body_font", value="Arial")]
    sec_styled_doc, sec_patch = apply_restyle(doc, sec_changes, section_ids=[target_sec.id])
    assert sec_styled_doc.sections[0].style.body_font == "Arial"
    assert sec_styled_doc.sections[1].style.body_font != "Arial"
    assert sec_patch.affected_unit_ids == [target_sec.id]


async def test_section_editor_agent():
    doc = sample_document()
    sec = doc.sections[0]
    updated_sec = sec.model_copy(update={"heading": "Updated Evidence Synthesis"})
    fake = FakeLLMClient({"section_editor": updated_sec})

    agent = SectionEditorAgent(fake)
    result = await agent.run(
        SectionEditorInput(
            section=sec,
            instruction="rename heading to Updated Evidence Synthesis",
            outline=[(s.id, s.heading) for s in doc.sections],
        )
    )
    assert result.heading == "Updated Evidence Synthesis"
    assert result.id == sec.id


async def test_section_insertion_agent():
    doc = sample_document()
    draft = SectionDraft(
        heading="Policy Implementation Roadmap",
        purpose="Action plan",
        level=1,
        blocks=[ParagraphBlock(text="Implementation phases 1 through 3.")],
    )
    fake = FakeLLMClient({"section_inserter": SectionInsertionPlan(sections=[draft])})
    agent = SectionInsertionAgent(fake)

    result = await agent.run(
        SectionInsertionInput(
            document=doc,
            instruction="Add an implementation roadmap section at the end",
            insertion_index=len(doc.sections),
            template_guidance="Follow default executive typography.",
        )
    )

    prompt = fake.calls_for("section_inserter")[0].messages[0].content
    assert result.sections[0].heading == "Policy Implementation Roadmap"
    assert f"Index: {len(doc.sections)}" in prompt
    assert doc.sections[-1].heading in prompt
    assert "Follow default executive typography" in prompt


async def test_style_planner_agent_document():
    doc = sample_document()
    plan = StylePlan(changes=[StyleChange(property="body_font", value="Calibri")])
    fake = FakeLLMClient({"style_planner": plan})
    agent = StylePlannerAgent(fake)

    result = await agent.run(
        StylePlannerInput(
            deck=doc,
            tokens=DesignTokens(),
            instruction="Change font to Calibri",
            scope_description="section sec_1",
            target_section_ids=[doc.sections[0].id],
        )
    )

    prompt = fake.calls_for("style_planner")[0].messages[0].content
    assert result.changes[0].value == "Calibri"
    assert doc.sections[0].id in prompt
    assert "section sec_1" in prompt


def test_docx_renderer_dynamic_styling(tmp_path):
    doc_plan = sample_document()
    custom_style = StyleSpec(
        heading_font="Georgia",
        body_font="Consolas",
        body_size_pt=12.0,
        margins_in=1.5,
        page_orientation="landscape",
        palette_overrides={"primary": "003366", "text": "222222"},
    )
    doc_plan = doc_plan.model_copy(update={"style": custom_style})

    out_docx = tmp_path / "styled_output.docx"
    render_result = render_document(doc_plan, DesignTokens(), out_docx)
    assert render_result.path.exists()

    rendered = Document(str(out_docx))
    section = rendered.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height
    assert abs(section.left_margin.inches - 1.5) < 0.01
    assert abs(section.top_margin.inches - 1.5) < 0.01
    assert rendered.styles["Normal"].font.name == "Consolas"
    assert rendered.styles["Heading 1"].font.name == "Georgia"
