from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from docforge.agents.workbook_planner import normalise
from docforge.design.system import generated_tokens
from docforge.ingest.datasets import load_dataset
from docforge.models.conversation import FileRole, Message, Role, Session, UploadedFile
from docforge.models.intent import Interpretation, OutputFormat, OutputSpec, ResearchNeed
from docforge.models.ir import ColumnSpec, SheetPlan, TableSpec, WorkbookPlan
from docforge.models.quality import Verdict
from docforge.orchestrator.generate import GenerationService
from docforge.validate.workbook import build_workbook
from tests.fixtures.workbooks import saas_model

pytestmark = pytest.mark.anyio

CSV = (
    b"Region,Sales,Margin,Date\nNorth,1200,12%,2026-01-31\n"
    b'South,"1,850.5",9%,2026-02-28\nEast,,15%,2026-03-31\n'
)


def test_dataset_inference_and_profile():
    ds = load_dataset("D1", "sales.csv", CSV)
    types = {c.name: c.dtype for c in ds.columns}
    assert types == {"Region": "text", "Sales": "float", "Margin": "percent", "Date": "date"}
    assert ds.rows[1][1] == 1850.5 and ds.rows[2][1] is None and ds.rows[0][2] == pytest.approx(0.12)
    profile = ds.profile()
    assert "Dataset D1 'sales.csv': 3 rows, 4 columns." in profile and "Sales (float; 1 missing" in profile


def test_dataset_from_xlsx_picks_largest_sheet():
    wb = Workbook()
    wb.active.append(["x"])
    big = wb.create_sheet("Data")
    big.append(["Name", "Count"])
    for i in range(5):
        big.append([f"n{i}", i])
    buf = io.BytesIO()
    wb.save(buf)
    ds = load_dataset("D1", "data.xlsx", buf.getvalue())
    assert [c.name for c in ds.columns] == ["Name", "Count"] and len(ds.rows) == 5


def test_normalise_sanitises_sheet_names_and_ids():
    plan = WorkbookPlan(
        title="t",
        sheets=[
            SheetPlan(id="x", name="Q1/Q2: [draft]", role="data"),
            SheetPlan(id="y", name="README", role="report"),
        ],
    )
    out = normalise(plan)
    assert [s.name for s in out.sheets] == ["Q1Q2 draft", "README 2"] and out.sheets[0].id == "sh-1"


def test_build_workbook_formulas_verified(tmp_path: Path):  # FR-XLSX-3/6
    build = build_workbook(
        saas_model(), generated_tokens("saas finance"), tmp_path, datasets={}, dossier=None, libreoffice=None
    )
    assert [f.message for f in build.report.findings] == []
    assert build.report.verdict == Verdict.pass_with_warnings  # only skipped tier-4/5 checks
    assert build.computed["REPORT!B1"] == pytest.approx(31525)
    assert build.computed["REVENUE MODEL!C2"] == pytest.approx(10500)
    wb = load_workbook(str(build.path))
    assert wb.sheetnames[0] == "README" and "Sources" not in wb.sheetnames
    assert wb["Revenue model"]["C3"].value == "=B3*(1+growth)"
    assert "growth" in wb.defined_names and wb["Assumptions"]["B3"].font.color.rgb.endswith("1F4FD1")
    assert wb["Revenue model"]._charts and wb["Revenue model"].tables


def test_formula_errors_and_expected_values_detected(tmp_path: Path):
    plan = saas_model(broken=True)
    plan.checks[0].expected = 99
    build = build_workbook(plan, generated_tokens("x"), tmp_path, datasets={}, dossier=None, libreoffice=None)
    found = {f.check_id: f.message for f in build.report.findings}
    assert "xlsx.formula_errors" in found and "C2=#" in found["xlsx.formula_errors"]  # unknown name
    assert "xlsx.expected_values" in found and build.report.verdict == Verdict.failed


def test_dataset_rows_fill_table(tmp_path: Path):
    ds = load_dataset("D1", "sales.csv", CSV)
    plan = WorkbookPlan(
        title="Sales",
        sheets=[
            SheetPlan(
                id="d",
                name="Data",
                role="data",
                tables=[
                    TableSpec(
                        name="Sales",
                        columns=[
                            ColumnSpec(name="Region"),
                            ColumnSpec(name="Sales", dtype="float"),
                            ColumnSpec(name="Margin", dtype="percent"),
                        ],
                        source_dataset="D1",
                    )
                ],
            )
        ],
    )
    build = build_workbook(
        plan, generated_tokens("x"), tmp_path, datasets={"D1": ds}, dossier=None, libreoffice=None
    )
    ws = load_workbook(str(build.path))["Data"]
    assert [c.value for c in ws[3]] == ["South", 1850.5, pytest.approx(0.09)]
    assert ws["C2"].number_format == "0.0%"
    assert not [f for f in build.report.findings if f.severity in ("critical", "major")]


async def test_generation_service_builds_and_revises_workbook(container, fake_llm):
    ref = container.blobs.put_bytes(CSV, ".csv")
    container.repos.sessions.put(Session(id="ses_1", owner="demo"))
    container.repos.files.put(
        UploadedFile(
            id="fil_d",
            session_id="ses_1",
            filename="sales.csv",
            media_type="text/csv",
            size=ref.size,
            sha256=ref.sha256,
            role=FileRole.dataset,
            blob_path=ref.path,
        )
    )
    fake_llm.add(
        "interpreter",
        Interpretation(
            kind="new_task",
            confidence=0.9,
            objective="SaaS model",
            output_spec=OutputSpec(format=OutputFormat.xlsx),
            research=ResearchNeed(depth="none"),
        ),
    )
    fake_llm.add(
        "workbook_planner", saas_model(broken=True), saas_model()
    )  # first plan fails, revision fixes it
    session = container.repos.sessions.require("ses_1")
    msg = container.repos.messages.put(Message(id="m1", session_id="ses_1", role=Role.user, content="model"))
    reply = await GenerationService(container).handle_message(session, msg)

    assert reply.artifact.format == OutputFormat.xlsx
    revise_prompt = fake_llm.calls_for("workbook_planner")[1].messages[-1].content
    assert "formula errors on 'Revenue model'" in revise_prompt
    assert "Dataset D1 'sales.csv'" in fake_llm.calls_for("workbook_planner")[0].messages[0].content
    assert (
        "Created an Excel workbook" in reply.message.content
        and "recalculated independently" in reply.message.content
    )
    assert reply.report.iterations == 1 and reply.report.verdict != Verdict.failed
    wb = load_workbook(io.BytesIO(container.blobs.read_bytes(reply.version.blob_path)))
    assert "Revenue model" in wb.sheetnames
