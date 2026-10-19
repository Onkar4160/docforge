from __future__ import annotations

import io
from pathlib import Path

import httpx
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from docforge.design.system import generated_tokens
from docforge.media.images import (
    ImageCandidate,
    ImageProvider,
    OpenverseProvider,
    PexelsProvider,
    build_image_provider,
    resolve_deck_images,
)
from docforge.models.ir import BulletItem, BulletsBlock, DeckPlan, ImageBlock, SlidePlan
from docforge.render.docx import render_document
from docforge.storage import BlobStore
from docforge.validate.deck import build_deck
from tests.conftest import make_settings
from tests.fixtures.documents import outline, written

pytestmark = pytest.mark.anyio


def _jpeg(w: int = 1600, h: int = 900) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (30, 90, 140)).save(buf, format="JPEG")
    return buf.getvalue()


class FakeImages(ImageProvider):
    name = "fake"

    def __init__(self, fail: bool = False) -> None:
        self.client = httpx.AsyncClient()
        self.queries: list[str] = []
        self.fail = fail

    async def search(self, query, *, count=5, wide=True):
        self.queries.append(query)
        return (
            []
            if self.fail
            else [
                ImageCandidate(
                    f"https://img/{len(self.queries)}.jpg",
                    1600,
                    900,
                    "Photo: A. Person (flickr), CC BY 2.0",
                    "https://flickr.com/p/1",
                )
            ]
        )

    async def download(self, candidate):
        return _jpeg()


def _deck() -> DeckPlan:
    return DeckPlan(
        title="EV charging",
        slides=[
            SlidePlan(
                id="s01",
                purpose="open",
                archetype="title",
                headline="EV charging in Europe",
                key_message="Scaling the network",
                blocks=[
                    ImageBlock(description="charging station at night", alt_text="Charging station at night")
                ],
            ),
            SlidePlan(
                id="s02",
                purpose="why",
                archetype="image_text",
                headline="Drivers want reliable fast charging",
                key_message="Reliability beats price",
                blocks=[
                    ImageBlock(description="driver plugging in electric car", alt_text="Driver plugging in"),
                    BulletsBlock(
                        items=[BulletItem(text="Uptime drives loyalty"), BulletItem(text="Speed matters")]
                    ),
                ],
            ),
            SlidePlan(
                id="s03",
                purpose="close",
                archetype="closing",
                headline="Next steps",
                blocks=[BulletsBlock(items=[BulletItem(text="Decide")])],
            ),
        ],
    )


def test_provider_selection(tmp_path: Path):
    assert isinstance(build_image_provider(make_settings(tmp_path, images_enabled=True)), OpenverseProvider)
    keyed = make_settings(tmp_path, images_enabled=True, pexels_api_key="px")
    assert isinstance(build_image_provider(keyed), PexelsProvider)
    assert build_image_provider(make_settings(tmp_path)) is None  # disabled in tests


async def test_openverse_mapping():
    def handler(req):
        assert req.url.params["license_type"] == "commercial,modification"
        return httpx.Response(
            200,
            json={
                "results": [
                    {
                        "url": "https://x/1.jpg",
                        "width": 1024,
                        "height": 683,
                        "license": "by",
                        "license_version": "2.0",
                        "creator": "NPS",
                        "provider": "flickr",
                        "foreign_landing_url": "https://flickr/1",
                    }
                ]
            },
        )

    p = OpenverseProvider(client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    [c] = await p.search("ev")
    assert c.credit == "Photo: NPS (flickr), CC BY 2.0" and c.landing_url == "https://flickr/1"


async def test_images_resolved_rendered_and_credited(tmp_path: Path):
    blobs = BlobStore(tmp_path / "blobs")
    deck = _deck()
    notes = await resolve_deck_images(deck, FakeImages(), blobs, limit=5)
    assert notes == []
    title_img, slide_img = deck.slides[0].blocks[0], deck.slides[1].blocks[0]
    assert title_img.asset_path and slide_img.credit.startswith("Photo: A. Person")

    build = build_deck(
        deck,
        generated_tokens("ev"),
        tmp_path / "out",
        spec=None,
        dossier=None,
        max_iterations=2,
        libreoffice=None,
        asset_resolver=blobs.resolve,
    )
    unresolved = [f.message for f in build.report.findings if f.status != "corrected"]
    assert unresolved == []
    prs = Presentation(str(build.path))
    names = [{s.name for s in slide.shapes} for slide in prs.slides]
    assert "content:hero-image" in names[0] and "content:image" in names[1]
    credits = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.name == "content:caption-credit"]
    assert credits == ["Photo: A. Person (flickr), CC BY 2.0"]
    pic = next(sh for sh in prs.slides[1].shapes if sh.name == "content:image")
    assert pic.crop_left > 0 or pic.crop_top > 0  # 16:9 photo cropped into a taller box


async def test_missing_images_fall_back_to_text_layouts(tmp_path: Path):
    deck = _deck()
    notes = await resolve_deck_images(deck, FakeImages(fail=True), BlobStore(tmp_path / "b"), limit=5)
    assert any("no suitable image" in n for n in notes)
    build = build_deck(
        deck,
        generated_tokens("ev"),
        tmp_path / "out",
        spec=None,
        dossier=None,
        max_iterations=1,
        libreoffice=None,
    )
    names = {s.name for s in Presentation(str(build.path)).slides[1].shapes}
    assert "content:image" not in names and "content:body" in names


async def test_document_image_with_caption(tmp_path: Path):
    from docforge.agents.doc_writer import assemble
    from docforge.media.images import resolve_document_images

    sections = written()
    sections[0].blocks.append(
        ImageBlock(description="water treatment plant", alt_text="A water treatment plant")
    )
    plan = assemble(outline(), sections)
    blobs = BlobStore(tmp_path / "blobs")
    await resolve_document_images(plan, FakeImages(), blobs, limit=2)
    render_document(plan, generated_tokens("x"), tmp_path / "d.docx", asset_resolver=blobs.resolve)
    doc = Document(str(tmp_path / "d.docx"))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Figure 2:")]
    assert captions == ["Figure 2: A water treatment plant. Photo: A. Person (flickr), CC BY 2.0."]
    assert len(doc.inline_shapes) == 2
