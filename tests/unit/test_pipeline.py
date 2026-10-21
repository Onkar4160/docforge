from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from docforge.agents.deck_planner import DeckPlannerAgent, PlannerInput, length_instruction, normalise
from docforge.agents.interpreter import InterpreterAgent, InterpreterInput, render_input
from docforge.core.errors import LLMOutputInvalid
from docforge.llm import FakeLLMClient
from docforge.models.conversation import Message, Role, Session, TaskStatus
from docforge.models.intent import (
    ClarificationQuestion,
    Interpretation,
    LengthSpec,
    OutputFormat,
    OutputSpec,
    RequestKind,
    ResearchNeed,
)
from docforge.models.ir import BulletItem, BulletsBlock, DeckPlan, SlidePlan
from docforge.orchestrator.generate import GenerationService
from tests.fixtures.decks import sample_deck

pytestmark = pytest.mark.anyio


def _interp(kind=RequestKind.new_task, fmt=OutputFormat.pptx, confidence=0.9, **kw) -> Interpretation:
    spec = OutputSpec(format=fmt, audience="board") if fmt else None
    return Interpretation(
        kind=kind,
        confidence=confidence,
        objective="EV charging board update",
        output_spec=spec,
        research=ResearchNeed(depth="none"),
        **kw,
    )


def _session(container) -> Session:
    return container.repos.sessions.put(Session(id="ses_1", owner="demo"))


def _msg(container, text: str) -> Message:
    return container.repos.messages.put(
        Message(id=f"msg_{abs(hash(text))}", session_id="ses_1", role=Role.user, content=text)
    )


# ---------------------------------------------------------------- agents (Stage 4 / 7 contracts)


def test_interpreter_input_rendering_includes_context():
    text = render_input(
        InterpreterInput(
            message="make slide 2 shorter",
            recent_turns=[("user", "deck please"), ("assistant", "done")],
            files=[("fil_1", "brand.pptx", "template")],
            active_task="Objective: x",
            artifact_outline=[("s02", "slide 2")],
        )
    )
    for fragment in (
        "<conversation>",
        "fil_1: brand.pptx (role: template)",
        "s02: slide 2",
        "<latest_user_message>\nmake slide 2 shorter",
    ):
        assert fragment in text


async def test_interpreter_drops_invented_file_ids():
    fake = FakeLLMClient({"interpreter": _interp(referenced_file_ids=["fil_1", "fil_ghost"])})
    result = await InterpreterAgent(fake).run(
        InterpreterInput(message="x", files=[("fil_1", "a.pdf", "source")])
    )
    assert result.referenced_file_ids == ["fil_1"]
    assert fake.requests[0].output_type is Interpretation


def test_planner_length_instruction_and_normalise():
    explicit = OutputSpec(format="pptx", length=LengthSpec(unit="slides", target=8, explicit=True))
    assert "exactly 8 slides" in length_instruction(explicit)
    assert "9-12" in length_instruction(OutputSpec(format="pptx"))
    deck = DeckPlan(
        title="t",
        design_hint="not-a-palette",
        slides=[
            SlidePlan(id="intro", purpose="p", archetype="title", headline="Hello."),
            SlidePlan(
                id="x9",
                purpose="p",
                archetype="bullets",
                headline="World",
                blocks=[BulletsBlock(items=[BulletItem(text="a")])],
            ),
        ],
    )
    out = normalise(deck)
    assert [s.id for s in out.slides] == ["s01", "s02"] and out.slides[0].headline == "Hello"
    assert out.design_hint is None


async def test_planner_revise_sends_previous_plan_and_problems():
    fake = FakeLLMClient({"deck_planner": sample_deck()})
    agent = DeckPlannerAgent(fake)
    inp = PlannerInput(objective="o", spec=OutputSpec(format="pptx"))
    first = await agent.run(inp)
    await agent.revise(inp, first, ["deck has 13 slides; 10 were requested"])
    last = fake.requests[-1]
    assert [m.role for m in last.messages] == ["user", "assistant", "user"]
    assert "10 were requested" in last.messages[-1].content
    assert "No researched evidence" in fake.requests[0].messages[0].content


# ---------------------------------------------------------------- generation service (M1)


async def test_generates_deck_and_persists_version(container, fake_llm):
    fake_llm.add("interpreter", _interp())
    fake_llm.add("deck_planner", sample_deck())
    session = _session(container)
    reply = await GenerationService(container).handle_message(
        session, _msg(container, "EV deck for the board")
    )

    assert reply.version is not None and reply.version.version == 1
    assert reply.report.verdict.value in ("pass", "pass_with_warnings")
    assert "13-slide deck" in reply.message.content and "Not checked" in reply.message.content
    assert reply.task.status == TaskStatus.delivered
    prs = Presentation(io.BytesIO(container.blobs.read_bytes(reply.version.blob_path)))
    assert len(prs.slides) == 13
    stored = container.repos.messages.list_by_parent("ses_1")
    assert [m.role for m in stored] == [Role.user, Role.assistant]
    assert stored[0].interpretation_id is not None
    assert (
        container.repos.quality_reports.require(reply.version.quality_report_id).artifact_id
        == reply.artifact.id
    )


async def test_regenerate_creates_version_two(container, fake_llm):
    fake_llm.add("interpreter", _interp(), _interp(kind=RequestKind.regenerate))
    fake_llm.add("deck_planner", sample_deck())
    session = _session(container)
    svc = GenerationService(container)
    first = await svc.handle_message(session, _msg(container, "deck"))
    second = await svc.handle_message(
        container.repos.sessions.require("ses_1"), _msg(container, "regenerate")
    )
    assert second.artifact.id == first.artifact.id and second.version.version == 2
    assert second.version.parent_version == 1


async def test_low_confidence_asks_instead_of_generating(container, fake_llm):
    fake_llm.add(
        "interpreter",
        _interp(
            confidence=0.3,
            clarification_questions=[
                ClarificationQuestion(question="Who is the audience?", options=["board", "engineers"])
            ],
        ),
    )
    reply = await GenerationService(container).handle_message(_session(container), _msg(container, "make it"))
    assert reply.version is None and "Who is the audience?" in reply.message.content
    assert reply.task.status == TaskStatus.awaiting_clarification
    assert fake_llm.calls_for("deck_planner") == []


async def test_plan_with_wrong_length_is_revised(container, fake_llm):
    interp = _interp()
    interp.output_spec.length = LengthSpec(unit="slides", target=13, explicit=True)
    short = sample_deck().model_copy(update={"slides": sample_deck().slides[:5]})
    fake_llm.add("interpreter", interp)
    fake_llm.add("deck_planner", short, sample_deck())
    reply = await GenerationService(container).handle_message(
        _session(container), _msg(container, "13 slides")
    )
    assert len(fake_llm.calls_for("deck_planner")) == 2
    assert len(reply.version.ir.slides) == 13


async def test_questions_are_answered_in_chat(container, fake_llm):
    from docforge.agents.answer import Answer

    fake_llm.add("interpreter", _interp(kind=RequestKind.question, fmt=None))
    fake_llm.add(
        "answer", Answer(text="I can build decks, documents and workbooks.", suggestions=["Make a deck"])
    )
    reply = await GenerationService(container).handle_message(_session(container), _msg(container, "what?"))
    assert reply.message.content.startswith("I can build decks") and "“Make a deck”" in reply.message.content
    assert reply.version is None


async def test_interpreter_failure_still_replies(container, fake_llm):
    from docforge.core.errors import LLMRateLimited

    def quota(_req):
        raise LLMRateLimited("all models out of quota", details={"quota": True})

    fake_llm.add("interpreter", quota)
    reply = await GenerationService(container).handle_message(
        _session(container), _msg(container, "try again")
    )
    assert "usage limit" in reply.message.content and "Nothing was delivered" in reply.message.content


async def test_unexpected_exception_still_replies(container, fake_llm):
    fake_llm.add("interpreter", _interp())

    def crash(_req):
        raise AttributeError("'PlaceholderPicture' object has no attribute 'text_frame'")

    fake_llm.add("deck_planner", crash)
    steps = []
    reply = await GenerationService(container).handle_message(
        _session(container), _msg(container, "x"), progress=lambda st, txt: steps.append(st)
    )
    assert "couldn't finish the presentation" in reply.message.content
    assert "AttributeError" in reply.message.content
    assert container.repos.tasks.require(reply.task.id).status == TaskStatus.failed
    assert steps == ["understand", "research", "plan"]


async def test_modify_request_restyles_deterministically(container, fake_llm):
    from docforge.models.intent import ChangeOperation, ChangeRequest, ChangeScope, ScopeLevel, StyleChange

    fake_llm.add("deck_planner", sample_deck())
    service = GenerationService(container)
    session = _session(container)
    fake_llm.add("interpreter", _interp())
    first = await service.handle_message(session, _msg(container, "deck"))
    modify = _interp(
        kind=RequestKind.modify_artifact,
        fmt=None,
        change=ChangeRequest(
            operation=ChangeOperation.restyle,
            scope=ChangeScope(level=ScopeLevel.style),
            instruction="change font to Georgia",
            style_changes=[StyleChange(property="heading_font", value="Georgia")],
        ),
    )
    fake_llm._queues["interpreter"].clear()
    fake_llm.add("interpreter", modify)
    reply = await service.handle_message(
        container.repos.sessions.require(session.id), _msg(container, "change font to Georgia")
    )
    assert first.version.version == 1 and reply.version is not None and reply.version.version == 2
    assert "Applied styling changes" in reply.message.content
    assert reply.version.ir.style.heading_font == "Georgia"


async def test_modify_request_edits_targeted_slide(container, fake_llm):
    from docforge.models.intent import (
        ChangeOperation,
        ChangeRequest,
        ChangeScope,
        ScopeLevel,
        TargetKind,
        TargetRef,
    )

    deck = sample_deck()
    fake_llm.add("deck_planner", deck)
    service = GenerationService(container)
    session = _session(container)
    fake_llm.add("interpreter", _interp())
    await service.handle_message(session, _msg(container, "deck"))

    # Prepare edited slide 2
    target_slide = deck.slides[1]
    edited_slide = target_slide.model_copy(update={"headline": "Brand New Headline for Slide 2"})
    fake_llm.add("slide_editor", edited_slide)

    modify = _interp(
        kind=RequestKind.modify_artifact,
        fmt=None,
        change=ChangeRequest(
            operation=ChangeOperation.rewrite,
            scope=ChangeScope(
                level=ScopeLevel.unit, targets=[TargetRef(kind=TargetKind.slide, locator="slide 2")]
            ),
            instruction="make slide 2 about market size",
        ),
    )
    fake_llm._queues["interpreter"].clear()
    fake_llm.add("interpreter", modify)
    reply = await service.handle_message(
        container.repos.sessions.require(session.id), _msg(container, "make slide 2 about market size")
    )
    assert reply.version is not None and reply.version.version == 2
    assert "Applied your change to slide “Brand New Headline for Slide 2”" in reply.message.content
    # Untouched slides remain identical
    assert reply.version.ir.slides[0] == deck.slides[0]
    assert reply.version.ir.slides[1].headline == "Brand New Headline for Slide 2"
    assert reply.version.ir.slides[2:] == deck.slides[2:]


async def test_modify_request_edits_targeted_document_section(container, fake_llm):
    from docforge.models.intent import (
        ChangeOperation,
        ChangeRequest,
        ChangeScope,
        ScopeLevel,
        TargetKind,
        TargetRef,
    )
    from tests.fixtures.documents import outline, written

    fake_llm.add("interpreter", _interp(fmt=OutputFormat.docx))
    fake_llm.add("doc_outliner", outline())
    for w in written():
        fake_llm.add("section_writer", w)

    service = GenerationService(container)
    session = _session(container)
    first = await service.handle_message(session, _msg(container, "report"))
    assert first.version.version == 1

    orig_doc = first.version.ir
    target_sec = orig_doc.sections[0]
    edited_sec = target_sec.model_copy(update={"heading": "Updated Evidence Section"})
    fake_llm.add("section_editor", edited_sec)

    modify = _interp(
        kind=RequestKind.modify_artifact,
        fmt=None,
        change=ChangeRequest(
            operation=ChangeOperation.rewrite,
            scope=ChangeScope(
                level=ScopeLevel.unit, targets=[TargetRef(kind=TargetKind.section, locator="section 1")]
            ),
            instruction="rename section 1",
        ),
    )
    fake_llm._queues["interpreter"].clear()
    fake_llm.add("interpreter", modify)
    reply = await service.handle_message(
        container.repos.sessions.require(session.id), _msg(container, "rename section 1")
    )
    assert reply.version is not None and reply.version.version == 2
    assert "Applied your change to section “Updated Evidence Section”" in reply.message.content
    assert reply.version.ir.sections[0].heading == "Updated Evidence Section"
    assert reply.version.ir.sections[1:] == orig_doc.sections[1:]


async def test_modify_request_edits_targeted_workbook_sheet(container, fake_llm):
    from docforge.models.intent import (
        ChangeOperation,
        ChangeRequest,
        ChangeScope,
        ScopeLevel,
        TargetKind,
        TargetRef,
    )
    from tests.fixtures.workbooks import saas_model

    fake_llm.add("interpreter", _interp(fmt=OutputFormat.xlsx))
    fake_llm.add("workbook_planner", saas_model())

    service = GenerationService(container)
    session = _session(container)
    first = await service.handle_message(session, _msg(container, "model"))
    assert first.version.version == 1

    orig_wb = first.version.ir
    target_sheet = orig_wb.sheets[0]
    edited_sheet = target_sheet.model_copy(update={"description": "Updated SaaS Assumptions"})
    fake_llm.add("sheet_editor", edited_sheet)

    modify = _interp(
        kind=RequestKind.modify_artifact,
        fmt=None,
        change=ChangeRequest(
            operation=ChangeOperation.rewrite,
            scope=ChangeScope(
                level=ScopeLevel.unit, targets=[TargetRef(kind=TargetKind.sheet, locator="sheet 1")]
            ),
            instruction="update description of assumptions",
        ),
    )
    fake_llm._queues["interpreter"].clear()
    fake_llm.add("interpreter", modify)
    reply = await service.handle_message(
        container.repos.sessions.require(session.id), _msg(container, "update description of assumptions")
    )
    assert reply.version is not None and reply.version.version == 2
    assert "Applied your change to sheet “Assumptions”" in reply.message.content
    assert reply.version.ir.sheets[0].description == "Updated SaaS Assumptions"
    assert reply.version.ir.sheets[1:] == orig_wb.sheets[1:]


async def test_planner_failure_marks_task_failed(container, fake_llm):
    fake_llm.add("interpreter", _interp())

    def boom(_req):
        raise LLMOutputInvalid("deck_planner: invalid structured output after repair")

    fake_llm.add("deck_planner", boom)
    reply = await GenerationService(container).handle_message(_session(container), _msg(container, "x"))
    assert "couldn't finish the presentation" in reply.message.content
    assert "unusable answer" in reply.message.content
    assert container.repos.tasks.require(reply.task.id).status == TaskStatus.failed


# ---------------------------------------------------------------- API end to end (AC-1 slice for S1)


def test_api_generate_download_and_report(auth_client: TestClient, fake_llm):
    fake_llm.add("interpreter", _interp())
    fake_llm.add("deck_planner", sample_deck())
    sid = auth_client.post("/api/sessions", json={}).json()["id"]
    res = auth_client.post(f"/api/sessions/{sid}/messages", json={"content": "Board deck on EV charging"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["reply"]["role"] == "assistant" and body["artifact"]["version"] == 1

    dl = auth_client.get(body["artifact"]["download_url"])
    assert dl.status_code == 200 and dl.content[:2] == b"PK"
    assert "attachment" in dl.headers["content-disposition"]
    assert len(Presentation(io.BytesIO(dl.content)).slides) == 13

    report = auth_client.get(body["artifact"]["quality_report_url"]).json()
    assert report["artifact_id"] == body["artifact"]["id"] and report["checks"]


def test_api_artifact_scoped_to_session(auth_client: TestClient, fake_llm):
    fake_llm.add("interpreter", _interp())
    fake_llm.add("deck_planner", sample_deck())
    sid = auth_client.post("/api/sessions", json={}).json()["id"]
    url = auth_client.post(f"/api/sessions/{sid}/messages", json={"content": "deck"}).json()["artifact"][
        "download_url"
    ]
    other = auth_client.post("/api/sessions", json={}).json()["id"]
    assert auth_client.get(url.replace(sid, other)).status_code == 404
