"""Programmatically built, licence-clean template fixtures (TESTING §4)."""

from __future__ import annotations

import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

SAMPLE_TITLE = "Sample Presentation Title Here"
SAMPLE_BULLET = "Replace this sample bullet text with your own content"


def _logo() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (240, 80), (192, 57, 43)).save(buf, format="PNG")
    return buf.getvalue()


def corporate_pptx(path: Path, *, widescreen: bool = False) -> Path:
    """Template with a custom theme (red accent, Georgia/Verdana) and a logo on every sample slide.

    Default is 4:3 (10 x 7.5 in); ``widescreen`` stretches masters and layouts to 13.33 x 7.5 in.
    """
    prs = Presentation()
    if widescreen:
        factor = 13.333 / 10
        prs.slide_width = Inches(13.333)
        for master in prs.slide_masters:
            # Layouts first: their placeholders may inherit geometry from the master.
            for shape in [*(s for lay in master.slide_layouts for s in lay.shapes), *master.shapes]:
                if shape.left is not None and shape.width is not None:
                    # Read all four first: setting one on an inherited placeholder zeroes the others.
                    x, y, w, h = shape.left, shape.top, shape.width, shape.height
                    shape.left, shape.top = int(x * factor), y
                    shape.width, shape.height = int(w * factor), h
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    for old, new in (
        ('<a:accent1>\n        <a:srgbClr val="4F81BD"/>', '<a:accent1>\n        <a:srgbClr val="C0392B"/>'),
    ):
        xml = xml.replace(old, new)
    import re

    xml = re.sub(r'(<a:accent1>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>C0392B", xml)
    xml = re.sub(r'(<a:accent2>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>2C3E50", xml)
    xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]+', r"\g<1>Georgia", xml)
    xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]+', r"\g<1>Verdana", xml)
    theme._blob = xml.encode("utf-8")

    logo = _logo()

    def add_logo(slide) -> None:
        pic = slide.shapes.add_picture(
            io.BytesIO(logo),
            prs.slide_width - Inches(1.4),
            prs.slide_height - Inches(0.6),
            Inches(1.2),
            Inches(0.4),
        )
        pic.name = "Company Logo"

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = SAMPLE_TITLE
    s1.placeholders[1].text = "Sample subtitle for the template preview"
    add_logo(s1)
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Sample Content Slide Heading"
    s2.placeholders[1].text = SAMPLE_BULLET
    add_logo(s2)
    s3 = prs.slides.add_slide(prs.slide_layouts[2])
    s3.shapes.title.text = "Sample Section Divider Heading"
    add_logo(s3)
    prs.save(str(path))
    return path
