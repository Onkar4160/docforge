"""Extract text from user-supplied sources (FR-RES-4) with locators (page/slide/sheet markers)."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from docforge.core.errors import UnsupportedFileError

MAX_CHARS = 400_000
MAX_SHEET_ROWS = 300


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[page {i}]\n{text}")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t_i, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        parts.append(f"[table {t_i}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [
            s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()
        ]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if texts or notes:
            slides.append(f"[slide {i}]\n" + "\n".join(texts) + (f"\n(notes) {notes}" if notes else ""))
    return "\n\n".join(slides)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i >= MAX_SHEET_ROWS:
                rows.append("… (truncated)")
                break
            if any(v is not None for v in row):
                rows.append("\t".join("" if v is None else str(v) for v in row))
        sheets.append(f"[sheet {ws.title}]\n" + "\n".join(rows))
    return "\n\n".join(sheets)


def _text(data: bytes, delimiter: str | None = None) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    if delimiter:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))[:MAX_SHEET_ROWS]
        return "\n".join("\t".join(r) for r in rows)
    return text


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    readers = {
        ".pdf": _pdf,
        ".docx": _docx,
        ".dotx": _docx,
        ".pptx": _pptx,
        ".potx": _pptx,
        ".xlsx": _xlsx,
        ".xltx": _xlsx,
        ".csv": lambda d: _text(d, ","),
        ".tsv": lambda d: _text(d, "\t"),
        ".txt": _text,
        ".md": _text,
    }
    if ext not in readers:
        raise UnsupportedFileError(f"cannot extract text from {ext or 'files without extension'}")
    return readers[ext](data)[:MAX_CHARS]
