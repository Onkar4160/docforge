"""Load user datasets (CSV/TSV/XLSX) with type inference and a compact profile (FR-XLSX-8).

Rows go into workbooks programmatically; only the profile goes to the LLM, so large datasets cost few tokens
and values are never retyped by a model.
"""

from __future__ import annotations

import csv
import io
import re
import statistics
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from docforge.core.errors import UnsupportedFileError

MAX_ROWS = 50_000
_NUM = re.compile(r"^-?[\d,]*\.?\d+(e-?\d+)?$", re.IGNORECASE)
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d.%m.%Y", "%Y-%m")


@dataclass
class ColumnProfile:
    name: str
    dtype: str  # int | float | percent | date | bool | text
    missing: int = 0
    minimum: float | None = None
    maximum: float | None = None
    mean: float | None = None
    distinct: int = 0
    examples: list[str] = field(default_factory=list)


@dataclass
class Dataset:
    id: str
    name: str
    columns: list[ColumnProfile]
    rows: list[list[object]]
    notes: list[str] = field(default_factory=list)

    def profile(self, max_examples: int = 3) -> str:
        lines = [f"Dataset {self.id} '{self.name}': {len(self.rows)} rows, {len(self.columns)} columns."]
        for c in self.columns:
            stats = ""
            if c.mean is not None:
                stats = f" min {c.minimum:g}, max {c.maximum:g}, mean {c.mean:.4g}"
            lines.append(
                f"- {c.name} ({c.dtype}; {c.missing} missing; {c.distinct} distinct;{stats}) "
                f"e.g. {', '.join(c.examples[:max_examples])}"
            )
        lines += [f"Note: {n}" for n in self.notes]
        return "\n".join(lines)


def _parse(value: object) -> tuple[object, str]:
    """Typed value + detected type for one raw cell."""
    if value is None:
        return None, "missing"
    if isinstance(value, bool):
        return value, "bool"
    if isinstance(value, int | float):
        return value, "int" if isinstance(value, int) or float(value).is_integer() else "float"
    if isinstance(value, datetime | date):
        return value, "date"
    text = str(value).strip()
    if not text or text.lower() in ("na", "n/a", "null", "none", "-", "nan"):
        return None, "missing"
    if text.lower() in ("true", "false", "yes", "no"):
        return text.lower() in ("true", "yes"), "bool"
    if text.endswith("%") and _NUM.match(text[:-1].replace(" ", "")):
        return float(text[:-1].replace(",", "")) / 100, "percent"
    stripped = text.replace(" ", "").lstrip("$€£¥")
    if _NUM.match(stripped):
        number = float(stripped.replace(",", ""))
        return (int(number) if number.is_integer() and "." not in stripped else number), (
            "int" if number.is_integer() and "." not in stripped else "float"
        )
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date(), "date"
        except ValueError:
            continue
    return text, "text"


def _column_type(kinds: list[str]) -> str:
    present = [k for k in kinds if k != "missing"]
    if not present:
        return "text"
    distinct = set(present)
    if distinct <= {"int"}:
        return "int"
    if distinct <= {"int", "float"}:
        return "float"
    if distinct <= {"percent"}:
        return "percent"
    for single in ("date", "bool"):
        if distinct == {single}:
            return single
    return "text"


def _raw_rows(filename: str, data: bytes) -> list[list[object]]:
    ext = Path(filename).suffix.lower()
    if ext in (".csv", ".tsv", ".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        delimiter = (
            "\t"
            if ext == ".tsv"
            else (csv.Sniffer().sniff(text[:4096], ",;\t|").delimiter if text.strip() else ",")
        )
        return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    if ext in (".xlsx", ".xltx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = max(wb.worksheets, key=lambda w: w.max_row or 0)
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raise UnsupportedFileError(f"datasets must be CSV, TSV or XLSX, not {ext}")


def load_dataset(dataset_id: str, filename: str, data: bytes) -> Dataset:
    raw = [r for r in _raw_rows(filename, data) if any(v not in (None, "") for v in r)]
    if not raw:
        raise UnsupportedFileError(f"{filename} contains no rows")
    header = [str(h).strip() if h not in (None, "") else f"Column {i + 1}" for i, h in enumerate(raw[0])]
    seen: dict[str, int] = {}
    for i, h in enumerate(header):  # unique, non-empty headers
        seen[h] = seen.get(h, 0) + 1
        if seen[h] > 1:
            header[i] = f"{h} ({seen[h]})"
    notes = []
    body = raw[1 : MAX_ROWS + 1]
    if len(raw) - 1 > MAX_ROWS:
        notes.append(f"truncated to the first {MAX_ROWS} rows")
    width = len(header)
    parsed = [[_parse(row[i] if i < len(row) else None) for i in range(width)] for row in body]
    columns, rows = [], [[v for v, _ in r] for r in parsed]
    for i, name in enumerate(header):
        kinds = [r[i][1] for r in parsed]
        dtype = _column_type(kinds)
        values = [r[i][0] for r in parsed]
        numeric = [float(v) for v in values if isinstance(v, int | float) and not isinstance(v, bool)]
        if dtype == "text":  # mixed columns: keep the original text, not partially parsed numbers
            for r_i, row in enumerate(rows):
                original = body[r_i][i] if i < len(body[r_i]) else None
                row[i] = None if values[r_i] is None else str(original).strip()
        prof = ColumnProfile(
            name=name,
            dtype=dtype,
            missing=sum(v is None for v in values),
            distinct=len({str(v) for v in values if v is not None}),
            examples=[str(v) for v in values if v is not None][:3],
        )
        if numeric and dtype in ("int", "float", "percent"):
            prof.minimum, prof.maximum, prof.mean = min(numeric), max(numeric), statistics.fmean(numeric)
        columns.append(prof)
    return Dataset(id=dataset_id, name=filename, columns=columns, rows=rows, notes=notes)
