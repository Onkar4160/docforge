"""Document plan checks, DOCX file checks and the build loop (QUALITY §2.1, §2.3)."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from docforge.core.ids import new_id
from docforge.models.design import DesignTokens
from docforge.models.evidence import ResearchDossier
from docforge.models.intent import LengthUnit, OutputSpec
from docforge.models.ir import (
    BulletsBlock,
    ChartBlock,
    DocumentPlan,
    KPIBlock,
    ParagraphBlock,
    TableBlock,
    block_claim_ids,
)
from docforge.models.quality import Dimension, FindingStatus, QualityReport, Severity
from docforge.render.docx import DocRenderResult, render_document
from docforge.validate.common import PLACEHOLDER_PATTERNS, CheckRun, similar
from docforge.validate.plan_checks import _FIGURE

WORDS_PER_PAGE = 450


def check_document_plan(run: CheckRun, plan: DocumentPlan, dossier: ResearchDossier | None) -> None:
    sections = plan.all_sections()
    with run.check("content.empty_units", 1):
        for s in sections:
            if not s.blocks and not s.subsections:
                run.finding(
                    "content.empty_units",
                    Dimension.content,
                    Severity.critical,
                    f"section '{s.heading}' has no content",
                    unit_id=s.id,
                )
    with run.check("content.duplicates", 1):
        for i, a in enumerate(sections):
            for b in sections[i + 1 :]:
                if similar(a.heading, b.heading) >= 0.92:
                    run.finding(
                        "content.duplicates",
                        Dimension.content,
                        Severity.major,
                        f"sections '{a.heading}' and '{b.heading}' duplicate each other",
                        unit_id=b.id,
                    )
    with run.check("evidence.citations_resolve", 1):
        known = {c.id for c in dossier.claims} if dossier else set()
        for s in sections:
            missing = sorted({cid for b in s.blocks for cid in block_claim_ids(b)} - known)
            if missing:
                run.finding(
                    "evidence.citations_resolve",
                    Dimension.evidence,
                    Severity.critical,
                    f"section '{s.heading}' cites unknown claims: {', '.join(missing[:5])}",
                    unit_id=s.id,
                )
    if dossier is not None and dossier.performed:
        with run.check("evidence.coverage", 1):
            for s in sections:
                for b in s.blocks:
                    texts = []
                    if isinstance(b, ParagraphBlock):
                        texts = [(b.text, b.claim_ids)]
                    elif isinstance(b, BulletsBlock):
                        texts = [(i.text, i.claim_ids or b.claim_ids) for i in b.items]
                    elif isinstance(b, ChartBlock | KPIBlock | TableBlock) and not block_claim_ids(b):
                        run.finding(
                            "evidence.coverage",
                            Dimension.evidence,
                            Severity.major,
                            f"data in section '{s.heading}' cites no evidence",
                            unit_id=s.id,
                        )
                    for text, cites in texts:
                        if _FIGURE.search(text) and not cites:
                            run.finding(
                                "evidence.coverage",
                                Dimension.evidence,
                                Severity.major,
                                f"uncited figure in '{s.heading}': '{text[:60]}…'",
                                unit_id=s.id,
                            )

def _iter_all_docx_paragraphs(doc):
    """Yield all paragraphs from body, table cells (including nested tables), and headers/footers."""
    yield from doc.paragraphs

    def _scan_table(tbl):
        for row in tbl.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from _scan_table(nested)

    for tbl in doc.tables:
        yield from _scan_table(tbl)

    for s in doc.sections:
        yield from s.header.paragraphs
        for tbl in s.header.tables:
            yield from _scan_table(tbl)
        yield from s.footer.paragraphs
        for tbl in s.footer.tables:
            yield from _scan_table(tbl)


def check_docx_file(
    run: CheckRun,
    path: Path,
    result: DocRenderResult,
    spec: OutputSpec | None,
    plan: DocumentPlan,
    libreoffice: Path | None,
) -> None:
    with run.check("common.reopen", 1):
        try:
            doc = Document(str(path))
        except Exception as exc:
            run.finding("common.reopen", Dimension.validity, Severity.critical, f"file does not open: {exc}")
            return
    body_xml = doc.element.body.xml

    with run.check("docx.heading_hierarchy", 1):
        previous = 0
        for level, text in result.headings:
            if level > previous + 1:
                run.finding(
                    "docx.heading_hierarchy",
                    Dimension.structure,
                    Severity.major,
                    f"heading '{text}' skips from level {previous} to {level}",
                )
            previous = level
        if sum(1 for p in doc.paragraphs if p.style.name == "Title") != 1:
            run.finding(
                "docx.heading_hierarchy",
                Dimension.structure,
                Severity.minor,
                "document should have exactly one Title paragraph",
            )

    with run.check("docx.toc_present", 1):
        if plan.front_matter.include_toc and "TOC \\o" not in body_xml:
            run.finding("docx.toc_present", Dimension.structure, Severity.minor, "table of contents missing")

    with run.check("docx.page_numbers", 1):
        if not any("PAGE" in s.footer._element.xml for s in doc.sections):
            run.finding(
                "docx.page_numbers", Dimension.structure, Severity.minor, "no page number field in footer"
            )

    with run.check("docx.references", 1):
        cited = set(re.findall(r"\[(\d+)(?:, \d+)*\]", "\n".join(p.text for p in doc.paragraphs)))
        has_refs = any(p.style.name.startswith("Heading") and p.text == "References" for p in doc.paragraphs)
        if (cited or result.references) and not has_refs:
            run.finding(
                "docx.references", Dimension.evidence, Severity.major, "citations without a references list"
            )

    with run.check("docx.tables", 1):
        for i, table in enumerate(doc.tables, 1):
            widths = {len(r.cells) for r in table.rows}
            if len(widths) > 1:
                run.finding("docx.tables", Dimension.structure, Severity.major, f"table {i} has ragged rows")

    with run.check("content.placeholder_text", 1):
        for p in _iter_all_docx_paragraphs(doc):
            if PLACEHOLDER_PATTERNS.search(p.text):
                run.finding(
                    "content.placeholder_text",
                    Dimension.content,
                    Severity.critical,
                    f"placeholder text: '{p.text[:60]}'",
                )
                break

    with run.check("common.length", 1):
        length = spec.length if spec else None
        if (
            length
            and length.explicit
            and length.target
            and length.unit in (LengthUnit.pages, LengthUnit.words)
        ):
            target_words = length.target * (WORDS_PER_PAGE if length.unit == LengthUnit.pages else 1)
            estimate = result.words + 120 * result.tables + 180 * result.figures
            if abs(estimate - target_words) > 0.15 * target_words:
                unit = length.unit.value
                actual = round(estimate / WORDS_PER_PAGE, 1) if unit == "pages" else estimate
                run.finding(
                    "common.length",
                    Dimension.structure,
                    Severity.major,
                    f"document is about {actual} {unit}; {length.target} were requested",
                    estimated=actual,
                    requested=length.target,
                )

    run.skipped(
        "docx.page_flow",
        4,
        "page flow needs a renderer (LibreOffice)"
        if libreoffice is None
        else "page flow analysis arrives with Stage 12",
    )
    if libreoffice is None:
        run.skipped("common.convert", 4, "LibreOffice not installed; conversion check not performed")
    run.skipped("content.review", 5, "LLM content review arrives with Stage 12")


@dataclass
class DocumentBuild:
    plan: DocumentPlan
    path: Path
    report: QualityReport
    result: DocRenderResult


def build_document(
    plan: DocumentPlan,
    tokens: DesignTokens,
    out_dir: Path,
    *,
    spec: OutputSpec | None,
    dossier: ResearchDossier | None,
    libreoffice: Path | None,
    template: Path | None = None,
    template_manifest=None,
    audience: str = "",
    artifact_id: str = "",
    version: int = 1,
    asset_resolver=None,
) -> DocumentBuild:
    path = out_dir / "document.docx"
    result = render_document(
        plan,
        tokens,
        path,
        dossier=dossier,
        template=template,
        template_manifest=template_manifest,
        audience=audience,
        citation_style=spec.citation_style if spec else "numeric",
        asset_resolver=asset_resolver,
    )
    run = CheckRun()
    check_document_plan(run, plan, dossier)
    check_docx_file(run, path, result, spec, plan, libreoffice)
    findings = [f.model_copy(update={"status": FindingStatus.unresolved}) for f in run.findings]
    report = QualityReport(
        id=new_id("report"),
        artifact_id=artifact_id,
        version=version,
        iterations=0,
        checks=run.checks,
        findings=findings,
    ).finalize()
    return DocumentBuild(plan, path, report, result)
