"""Inspect a rendered PPTX file (QUALITY §2.2). Reads the file, not the plan, so it judges what ships."""

from __future__ import annotations

import math
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from docforge.design.system import contrast_ratio
from docforge.layout.backdrop import color_under
from docforge.layout.text import EMU_PER_INCH, Para, measure
from docforge.models.quality import Dimension, Severity
from docforge.models.template import ElementKind, ElementRole, TemplateManifest
from docforge.validate.common import PLACEHOLDER_PATTERNS, CheckRun

TOLERANCE = int(0.02 * EMU_PER_INCH)
_FLOORS = (
    ("content:headline", 20),
    ("content:title", 24),
    ("footer:", 9),
    ("content:caption", 10),
    ("content:references", 10),
    ("content:audience", 10),
    ("content:", 14),
)
_DENSITY_WORDS = 90


def _floor(name: str) -> float | None:
    return next((size for prefix, size in _FLOORS if name.startswith(prefix)), None)


def _box(shape) -> tuple[int, int, int, int]:
    return int(shape.left or 0), int(shape.top or 0), int(shape.width or 0), int(shape.height or 0)


def _intersection(a, b) -> int:
    dx = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    dy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return dx * dy if dx > 0 and dy > 0 else 0


def _paras(shape, default_font: str | None = None) -> tuple[list[Para], list[tuple[float, bool, str | None]]]:
    """Measurement paragraphs plus per-paragraph (size, bold, colour) for typography checks."""
    paras, styles = [], []
    for p in shape.text_frame.paragraphs:
        runs = p.runs
        text = "".join(r.text for r in runs)
        first = runs[0].font if runs else None
        size = first.size.pt if first is not None and first.size is not None else 18.0
        bold = bool(first.bold) if first is not None else False
        color = None
        if first is not None and first.color is not None and first.color.type is not None:
            try:
                color = str(first.color.rgb)
            except AttributeError:
                color = None
        p_pr = p._p.pPr
        indent = int(p_pr.get("marL", "0")) if p_pr is not None else 0
        after = p.space_after.pt if p.space_after is not None else 0.0
        before = p.space_before.pt if p.space_before is not None else 0.0
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        paras.append(
            Para(
                text,
                size,
                bold=bold,
                font=(first.name if first is not None else None) or default_font,
                indent_emu=indent,
                space_before_pt=before,
                space_after_pt=after,
                line_spacing=spacing,
            )
        )
        styles.append((size, bold, color))
    return paras, styles


def _solid(fill) -> str | None:
    try:
        if fill.type == MSO_FILL.SOLID:
            return str(fill.fore_color.rgb)
    except (AttributeError, TypeError):
        return None
    return None


def check_pptx_file(
    run: CheckRun,
    path: Path,
    unit_ids: list[str],
    *,
    libreoffice: Path | None,
    theme_fonts: tuple[str | None, str | None] = (None, None),
) -> None:
    with run.check("common.reopen", 1):
        try:
            prs = Presentation(str(path))
        except Exception as exc:
            run.finding("common.reopen", Dimension.validity, Severity.critical, f"file does not open: {exc}")
            return

    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    scale = slide_h / (7.5 * EMU_PER_INCH)  # minimum sizes are defined for a 7.5-inch-tall slide
    major, minor = theme_fonts

    def default_font(shape) -> str | None:
        return major if shape.is_placeholder and shape.placeholder_format.idx == 0 else minor

    slides = list(prs.slides)

    def unit(i: int) -> str | None:
        return unit_ids[i] if i < len(unit_ids) else None

    with run.check("pptx.out_of_bounds", 2):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                if not shape.name.startswith(("content:", "footer:", "decor:", "branding:")):
                    continue
                x, y, w, h = _box(shape)
                if (
                    x < -TOLERANCE
                    or y < -TOLERANCE
                    or x + w > slide_w + TOLERANCE
                    or y + h > slide_h + TOLERANCE
                ):
                    run.finding(
                        "pptx.out_of_bounds",
                        Dimension.layout,
                        Severity.critical,
                        f"'{shape.name}' extends beyond the slide",
                        unit_id=unit(i),
                        element=shape.name,
                    )

    with run.check("pptx.text_overflow", 2):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                if not shape.has_text_frame or not shape.text_frame.text.strip():
                    continue
                tf = shape.text_frame
                if tf.word_wrap is False:  # badges: single short label, not wrapped
                    continue
                paras, _ = _paras(shape, default_font(shape))
                width = int(shape.width) - int(tf.margin_left) - int(tf.margin_right)
                height = int(shape.height) - int(tf.margin_top) - int(tf.margin_bottom)
                m = measure(paras, width)
                if m.height_emu > height * 1.02 or m.widest_word_overflow:
                    over = (m.height_emu - height) / EMU_PER_INCH
                    run.finding(
                        "pptx.text_overflow",
                        Dimension.layout,
                        Severity.critical,
                        f"text in '{shape.name}' overflows its box by {max(over, 0):.2f} in"
                        + (" (a word is wider than the box)" if m.widest_word_overflow else ""),
                        unit_id=unit(i),
                        element=shape.name,
                        estimated=m.estimated,
                    )

    with run.check("pptx.font_floor", 2):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                floor = _floor(shape.name)
                floor = round(floor * scale, 1) if floor is not None else None
                if floor is None or not shape.has_text_frame:
                    continue
                _, styles = _paras(shape)
                small = [s for s, _, _ in styles if s < floor - 0.05]
                if small:
                    run.finding(
                        "pptx.font_floor",
                        Dimension.typography,
                        Severity.major,
                        f"'{shape.name}' uses {min(small):.0f}pt, below the {floor}pt minimum",
                        unit_id=unit(i),
                        element=shape.name,
                    )

    with run.check("pptx.overlap", 2):
        for i, slide in enumerate(slides):
            content = [s for s in slide.shapes if s.name.startswith("content:")]
            for a_i, a in enumerate(content):
                for b in content[a_i + 1 :]:
                    inter = _intersection(_box(a), _box(b))
                    smaller = min(_box(a)[2] * _box(a)[3], _box(b)[2] * _box(b)[3]) or 1
                    if inter / smaller > 0.02:
                        run.finding(
                            "pptx.overlap",
                            Dimension.layout,
                            Severity.major,
                            f"'{a.name}' overlaps '{b.name}'",
                            unit_id=unit(i),
                            element=a.name,
                        )

    with run.check("pptx.empty_placeholder", 1):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                    run.finding(
                        "pptx.empty_placeholder",
                        Dimension.layout,
                        Severity.critical,
                        f"empty placeholder '{shape.name}'",
                        unit_id=unit(i),
                        element=shape.name,
                    )
                elif (
                    shape.name.startswith("content:")
                    and shape.has_text_frame
                    and not shape.text_frame.text.strip()
                ):
                    run.finding(
                        "pptx.empty_placeholder",
                        Dimension.layout,
                        Severity.major,
                        f"empty text box '{shape.name}'",
                        unit_id=unit(i),
                        element=shape.name,
                    )

    with run.check("content.placeholder_text", 1):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                if shape.has_text_frame and PLACEHOLDER_PATTERNS.search(shape.text_frame.text):
                    run.finding(
                        "content.placeholder_text",
                        Dimension.content,
                        Severity.critical,
                        f"placeholder text in '{shape.name}'",
                        unit_id=unit(i),
                        element=shape.name,
                    )

    with run.check("pptx.chart_data", 1):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                if not getattr(shape, "has_chart", False):
                    continue
                chart = shape.chart
                n_cat = len(list(chart.plots[0].categories))
                for plot in chart.plots:
                    for series in plot.series:
                        values = list(series.values)
                        if len(values) != n_cat or any(v is None or math.isnan(v) for v in values):
                            run.finding(
                                "pptx.chart_data",
                                Dimension.data,
                                Severity.critical,
                                f"chart series '{series.name}' does not match its categories",
                                unit_id=unit(i),
                                element=shape.name,
                            )

    with run.check("pptx.density", 2):
        for i, slide in enumerate(slides):
            words = sum(
                len(s.text_frame.text.split())
                for s in slide.shapes
                if s.has_text_frame and s.name.startswith("content:") and s.name != "content:headline"
            )
            if words > _DENSITY_WORDS:
                run.finding(
                    "pptx.density",
                    Dimension.content,
                    Severity.minor,
                    f"slide carries {words} words of body text (aim for ≤ {_DENSITY_WORDS})",
                    unit_id=unit(i),
                    words=words,
                )

    with run.check("a11y.alt_text", 1):
        for i, slide in enumerate(slides):
            for shape in slide.shapes:
                ours = shape.name.startswith(("content:", "branding:"))
                if ours and (
                    getattr(shape, "has_chart", False) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                ):
                    c_nv_pr = shape._element.xpath("./*[1]/p:cNvPr")
                    if not c_nv_pr or not c_nv_pr[0].get("descr"):
                        run.finding(
                            "a11y.alt_text",
                            Dimension.accessibility,
                            Severity.minor,
                            f"'{shape.name}' has no alternative text",
                            unit_id=unit(i),
                            element=shape.name,
                        )

    with run.check("a11y.contrast", 2):
        for i, slide in enumerate(slides):
            background = _solid(slide.background.fill) or "FFFFFF"
            shapes = list(slide.shapes)
            for shape in shapes:
                if not shape.name.startswith(("content:", "footer:")) or not shape.has_text_frame:
                    continue
                under = color_under(slide, _box(shape), stop_at=shape._element, default=background)
                for size, bold, color in _paras(shape)[1]:
                    if not color:
                        continue
                    needed = 3.0 if size >= 18 or (bold and size >= 14) else 4.5
                    ratio = contrast_ratio(color, under)
                    if ratio < needed:
                        run.finding(
                            "a11y.contrast",
                            Dimension.accessibility,
                            Severity.major,
                            f"text in '{shape.name}' has contrast {ratio:.1f}:1 (needs {needed}:1)",
                            unit_id=unit(i),
                            element=shape.name,
                        )
                        break

    if libreoffice is None:
        run.skipped("common.convert", 4, "LibreOffice not installed; conversion check not performed")
    else:
        with run.check("common.convert", 4), tempfile.TemporaryDirectory() as tmp:
            proc = subprocess.run(  # noqa: S603 - fixed executable path from configuration
                [str(libreoffice), "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
            if proc.returncode != 0 or not any(Path(tmp).glob("*.pdf")):
                run.finding(
                    "common.convert",
                    Dimension.validity,
                    Severity.critical,
                    "LibreOffice could not convert the deck",
                    stderr=proc.stderr[-300:],
                )
    run.skipped("pptx.visual_clipping", 4, "pixel-level clipping comparison arrives with Stage 12")
    run.skipped("content.review", 5, "LLM content review arrives with Stage 7")


def check_template_fidelity(
    run: CheckRun,
    path: Path,
    manifest: TemplateManifest,
    *,
    allowed_fonts: set[str] | None = None,
) -> None:
    """Template identity kept, sample residue gone (QUALITY §2.2: template checks)."""
    prs = Presentation(str(path))
    theme_fonts = {f for f in (manifest.theme.major_font, manifest.theme.minor_font) if f}
    theme_fonts |= set(manifest.pptx.slide_fonts.values())
    theme_fonts |= {
        e.font.name
        for inv in manifest.pptx.sample_slides
        for e in inv.elements
        if e.font is not None and e.font.name
    }
    theme_fonts |= allowed_fonts or set()
    slides_mode = manifest.pptx.design_mode == "slides"
    samples = {
        e.text
        for inv in manifest.pptx.sample_slides
        for e in inv.elements
        if e.role == ElementRole.sample_content and e.text and len(e.text) > 12
    }
    branding = sum(
        1
        for inv in manifest.pptx.sample_slides[:1]
        for e in inv.elements
        if e.role == ElementRole.branding and e.kind == ElementKind.picture
    )
    with run.check("pptx.template_fonts", 1):
        foreign = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    foreign |= {
                        r.font.name
                        for p in shape.text_frame.paragraphs
                        for r in p.runs
                        if r.font.name and r.font.name not in theme_fonts
                    }
        if theme_fonts and foreign:
            run.finding(
                "pptx.template_fonts",
                Dimension.template,
                Severity.major,
                f"fonts outside the template theme: {', '.join(sorted(foreign))}",
            )
    with run.check("pptx.sample_residue", 1):
        texts = [s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame]
        residue = [t for t in samples if any(t in x for x in texts)]
        if residue:
            run.finding(
                "pptx.sample_residue",
                Dimension.template,
                Severity.critical,
                f"template sample text remains: '{residue[0][:60]}'",
            )
    with run.check("pptx.branding_retained", 1):
        if branding and not slides_mode:  # clone mode copies the artwork itself
            for i, slide in enumerate(prs.slides):
                present = sum(1 for s in slide.shapes if s.name.startswith("branding:"))
                if present < branding:
                    run.finding(
                        "pptx.branding_retained",
                        Dimension.template,
                        Severity.major,
                        f"slide {i + 1} lost template branding",
                        element=f"slide {i + 1}",
                    )
