"""Workbook checks and build (QUALITY §2.4, ADR-011).

Formulas are verified by recalculating the saved file with an independent engine (``formulas``), not by
trusting the plan: every error value, failed expected-value check and missing formula is reported.
"""

from __future__ import annotations

import contextlib
import os
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

from docforge.core.ids import new_id
from docforge.ingest.datasets import Dataset
from docforge.models.design import DesignTokens
from docforge.models.evidence import ResearchDossier
from docforge.models.ir import WorkbookPlan
from docforge.models.quality import Dimension, FindingStatus, QualityReport, Severity
from docforge.render.xlsx import WorkbookRenderResult, render_workbook
from docforge.validate.common import CheckRun

ERROR_VALUES = {"#DIV/0!", "#NAME?", "#REF!", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}
_KEY = re.compile(r"^'\[[^\]]+\](.+)'!([A-Z]+[0-9]+)$")


def recalculate(path: Path) -> dict[str, object]:
    """Compute every formula; returns {'SHEET!A1': value} with upper-case sheet names."""
    os.environ.setdefault("TQDM_DISABLE", "1")
    import formulas  # heavy import, only when validating workbooks

    model = formulas.ExcelModel().loads(str(path)).finish()
    out: dict[str, object] = {}
    for key, rng in model.calculate().items():
        match = _KEY.match(key)
        if match is None:
            continue
        value = getattr(rng, "value", rng)
        with contextlib.suppress(TypeError, IndexError, KeyError):
            value = value[0][0]  # single cells come back as 1x1 arrays
        out[f"{match.group(1).upper()}!{match.group(2)}"] = value
    return out


def _is_error(value: object) -> bool:
    return str(value) in ERROR_VALUES


def check_workbook_plan(
    run: CheckRun, plan: WorkbookPlan, dossier: ResearchDossier | None, datasets: dict[str, Dataset]
) -> None:
    with run.check("evidence.citations_resolve", 1):
        known = {c.id for c in dossier.claims} if dossier else set()
        for s in plan.sheets:
            missing = sorted(set(s.claim_ids) - known)
            if missing:
                run.finding(
                    "evidence.citations_resolve",
                    Dimension.evidence,
                    Severity.critical,
                    f"sheet '{s.name}' cites unknown claims: {', '.join(missing[:5])}",
                    unit_id=s.id,
                )
    with run.check("content.empty_units", 1):
        for s in plan.sheets:
            if not s.tables and not s.cells:
                run.finding(
                    "content.empty_units",
                    Dimension.content,
                    Severity.critical,
                    f"sheet '{s.name}' has no content",
                    unit_id=s.id,
                )
            for t in s.tables:
                if t.source_dataset and t.source_dataset not in datasets:
                    run.finding(
                        "content.empty_units",
                        Dimension.data,
                        Severity.critical,
                        f"table '{t.name}' refers to unknown dataset {t.source_dataset}",
                        unit_id=s.id,
                    )
                elif not t.rows and not t.source_dataset:
                    run.finding(
                        "content.empty_units",
                        Dimension.data,
                        Severity.major,
                        f"table '{t.name}' on '{s.name}' has no rows",
                        unit_id=s.id,
                    )


def check_xlsx_file(
    run: CheckRun, path: Path, plan: WorkbookPlan, result: WorkbookRenderResult, libreoffice: Path | None
) -> dict[str, object]:
    computed: dict[str, object] = {}
    with run.check("common.reopen", 1):
        try:
            wb = load_workbook(str(path))
        except Exception as exc:
            run.finding("common.reopen", Dimension.validity, Severity.critical, f"file does not open: {exc}")
            return computed
    sheet_ids = {s.name: s.id for s in plan.sheets}

    with run.check("xlsx.hardcoded_calcs", 1):
        literal = []
        for ref in result.calculated_cells:
            sheet, cell = ref.rsplit("!", 1)
            value = wb[sheet][cell].value
            if not (isinstance(value, str) and value.startswith("=")):
                literal.append(ref)
        if literal:
            run.finding(
                "xlsx.hardcoded_calcs",
                Dimension.data,
                Severity.major,
                f"{len(literal)} calculated cells hold values instead of formulas (e.g. {literal[0]})",
            )

    with run.check("xlsx.number_formats", 1):
        wrong = [
            ref
            for ref, fmt in result.typed_cells
            if wb[ref.rsplit("!", 1)[0]][ref.rsplit("!", 1)[1]].number_format != fmt
        ]
        if wrong:
            run.finding(
                "xlsx.number_formats",
                Dimension.data,
                Severity.minor,
                f"{len(wrong)} typed cells lost their number format (e.g. {wrong[0]})",
            )

    with run.check("xlsx.chart_ranges", 1):
        for s in plan.sheets:
            ws = wb[s.name]
            for chart in s.charts:
                min_col, min_row, max_col, max_row = range_boundaries(chart.data_ref)
                empty = sum(
                    1
                    for row in ws.iter_rows(
                        min_row=min_row + 1, max_row=max_row, min_col=min_col, max_col=max_col
                    )
                    for c in row
                    if c.value is None
                )
                if empty:
                    run.finding(
                        "xlsx.chart_ranges",
                        Dimension.data,
                        Severity.major,
                        f"chart '{chart.title}' on '{s.name}' covers {empty} empty cells",
                        unit_id=s.id,
                    )

    with run.check("xlsx.sources_sheet", 1):
        if any(s.claim_ids for s in plan.sheets) and "Sources" not in wb.sheetnames:
            run.finding(
                "xlsx.sources_sheet",
                Dimension.evidence,
                Severity.major,
                "figures cite evidence but the workbook has no Sources sheet",
            )

    with run.check("xlsx.formula_errors", 3):
        computed = recalculate(path)
        errors = sorted(k for k, v in computed.items() if _is_error(v))
        by_sheet: dict[str, list[str]] = {}
        for key in errors:
            by_sheet.setdefault(key.split("!")[0], []).append(f"{key.split('!')[1]}={computed[key]}")
        for sheet_upper, cells in by_sheet.items():
            name = next((n for n in sheet_ids if n.upper() == sheet_upper), sheet_upper)
            run.finding(
                "xlsx.formula_errors",
                Dimension.data,
                Severity.critical,
                f"formula errors on '{name}': {', '.join(cells[:5])}"
                + (f" (+{len(cells) - 5} more)" if len(cells) > 5 else ""),
                unit_id=sheet_ids.get(name),
                cells=cells[:20],
            )

    with run.check("xlsx.expected_values", 3):
        names = {s.id: s.name for s in plan.sheets}
        for chk in plan.checks:
            key = f"{names.get(chk.sheet_id, '?').upper()}!{chk.cell.replace('$', '')}"
            actual = computed.get(key)
            try:
                ok = actual is not None and abs(float(actual) - chk.expected) <= max(
                    chk.tolerance, abs(chk.expected) * 1e-6
                )
            except (TypeError, ValueError):
                ok = False
            if not ok:
                run.finding(
                    "xlsx.expected_values",
                    Dimension.data,
                    Severity.critical,
                    f"{names.get(chk.sheet_id)}!{chk.cell} = {actual}, expected {chk.expected:g}"
                    + (f" ({chk.rationale})" if chk.rationale else ""),
                    unit_id=chk.sheet_id,
                )

    if libreoffice is None:
        run.skipped("common.convert", 4, "LibreOffice not installed; conversion check not performed")
    run.skipped("content.review", 5, "LLM content review arrives with Stage 12")
    return computed


@dataclass
class WorkbookBuild:
    plan: WorkbookPlan
    path: Path
    report: QualityReport
    result: WorkbookRenderResult
    computed: dict[str, object]


def build_workbook(
    plan: WorkbookPlan,
    tokens: DesignTokens,
    out_dir: Path,
    *,
    datasets: dict[str, Dataset],
    dossier: ResearchDossier | None,
    libreoffice: Path | None,
    artifact_id: str = "",
    version: int = 1,
) -> WorkbookBuild:
    path = out_dir / "workbook.xlsx"
    result = render_workbook(plan, tokens, path, datasets=datasets, dossier=dossier)
    run = CheckRun()
    check_workbook_plan(run, plan, dossier, datasets)
    computed = check_xlsx_file(run, path, plan, result, libreoffice)
    findings = [f.model_copy(update={"status": FindingStatus.unresolved}) for f in run.findings]
    report = QualityReport(
        id=new_id("report"),
        artifact_id=artifact_id,
        version=version,
        iterations=0,
        checks=run.checks,
        findings=findings,
    ).finalize()
    return WorkbookBuild(plan, path, report, result, computed)


def extract_workbook_diagnostics(plan: WorkbookPlan, build: WorkbookBuild) -> list[str]:
    """Extract actionable formula, calculation, and constraint diagnostics for LLM revision."""
    diagnostics: list[str] = []
    named_ranges_desc = ", ".join(f"'{nr.name}' ({nr.ref})" for nr in plan.named_ranges) or "none"
    sheet_names = [s.name for s in plan.sheets]

    # 1. Formula errors from computed cells
    errors = {k: v for k, v in build.computed.items() if _is_error(v)}
    for key, err_val in sorted(errors.items())[:12]:
        sheet_part, cell_coord = key.split("!", 1)
        sheet = next((s for s in plan.sheets if s.name.upper() == sheet_part), None)
        sheet_name = sheet.name if sheet else sheet_part

        # Find formula in plan
        formula = None
        if sheet:
            for c in sheet.cells:
                if c.ref.upper() == cell_coord:
                    formula = c.formula
                    break
            if not formula:
                for t in sheet.tables:
                    try:
                        col0_letter, row0 = coordinate_from_string(t.anchor)
                        col0 = column_index_from_string(col0_letter)
                        target_col_letter, target_row = coordinate_from_string(cell_coord)
                        target_col = column_index_from_string(target_col_letter)
                        col_idx = target_col - col0
                        if 0 <= col_idx < len(t.columns):
                            col = t.columns[col_idx]
                            if col.formula_template:
                                formula = col.formula_template.format(row=target_row)
                    except Exception:
                        pass

        formula_str = f"'{formula}'" if formula else "unknown formula"
        hint = ""
        if str(err_val) == "#NAME?":
            hint = (
                f"Check function name spelling or named range. Available named ranges: {named_ranges_desc}."
            )
        elif str(err_val) == "#REF!":
            hint = f"Cell or sheet reference does not exist. Sheets: {', '.join(sheet_names)}."
        elif str(err_val) == "#DIV/0!":
            hint = "Division by zero or empty cell; wrap with IFERROR, e.g. =IFERROR(..., 0)."
        elif str(err_val) == "#VALUE!":
            hint = "Data type mismatch in operation (e.g. arithmetic on text)."

        msg = (
            f"formula errors on '{sheet_name}': Cell {cell_coord} with formula "
            f"{formula_str} evaluated to {err_val}. {hint}"
        ).strip()
        diagnostics.append(msg)

    # 2. Hardcoded calculations and expected value checks
    for f in build.report.findings:
        if (
            f.check_id in ("xlsx.hardcoded_calcs", "xlsx.expected_values", "xlsx.chart_ranges")
            and f.severity in (Severity.critical, Severity.major)
            and f.message not in diagnostics
        ):
            diagnostics.append(f.message)

    return diagnostics
