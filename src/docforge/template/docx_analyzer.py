"""Deterministic DOCX template inspection into the shared TemplateManifest."""

from __future__ import annotations

import hashlib
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from docforge.models.intent import OutputFormat
from docforge.models.template import (
    DocxTemplateInfo,
    FontSpec,
    HeaderFooterElementInfo,
    HeaderFooterInfo,
    ImageAssetInfo,
    PropertySource,
    SectionInfo,
    StyleInfo,
    TemplateManifest,
    ThemeSpec,
)

ANALYZER_VERSION = 3


def _font(style) -> FontSpec | None:
    font = getattr(style, "font", None)
    if font is None:
        return None
    base_style = getattr(style, "base_style", None)
    base = getattr(base_style, "font", None) if base_style is not None else None
    values = {}
    sources = {}
    for name in ("name", "size", "bold"):
        value = getattr(font, name)
        inherited = getattr(base, name) if base is not None else None
        values[name] = value if value is not None else inherited
        sources[name] = (
            PropertySource.explicit
            if value is not None
            else PropertySource.inherited
            if inherited is not None
            else PropertySource.unknown
        )
    color = None
    color_source = PropertySource.unknown
    if font.color is not None:
        if font.color.rgb is not None:
            color = str(font.color.rgb).upper()
            color_source = PropertySource.explicit
        elif font.color.theme_color is not None:
            color = f"theme:{font.color.theme_color.name.lower().replace('_', '')}"
            color_source = PropertySource.explicit
    if color is None and base is not None and base.color is not None:
        if base.color.rgb is not None:
            color = str(base.color.rgb).upper()
            color_source = PropertySource.inherited
        elif base.color.theme_color is not None:
            color = f"theme:{base.color.theme_color.name.lower().replace('_', '')}"
            color_source = PropertySource.inherited
    if not any(value is not None for value in values.values()) and color is None:
        return None
    return FontSpec(
        name=values["name"],
        size_pt=values["size"].pt if values["size"] is not None else None,
        bold=values["bold"],
        color=color,
        name_source=sources["name"],
        size_source=sources["size"],
        bold_source=sources["bold"],
        color_source=color_source,
    )


def _extract_theme_colors(path: Path) -> dict[str, str]:
    colors: dict[str, str] = {}
    try:
        with zipfile.ZipFile(str(path)) as z:
            if "word/theme/theme1.xml" in z.namelist():
                tree = ET.fromstring(z.read("word/theme/theme1.xml"))
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for child in tree.findall(".//a:clrScheme/*", ns):
                    tag = child.tag.split("}")[-1]
                    srgb = child.find(".//a:srgbClr", ns)
                    if srgb is not None and srgb.get("val"):
                        colors[tag] = srgb.get("val").upper()
                    else:
                        sys = child.find(".//a:sysClr", ns)
                        if sys is not None and sys.get("lastClr"):
                            colors[tag] = sys.get("lastClr").upper()
    except Exception:
        pass
    return colors


def _extract_image_asset(
    part,
    drawing_elem,
    location: str,
    section_index: int,
    image_assets: list[ImageAssetInfo],
    seen_hashes: dict[str, str],
    cache_dir: Path | None,
) -> tuple[str | None, int | None, int | None, str]:
    for blip in drawing_elem.xpath(".//a:blip"):
        rid = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        if not rid:
            continue
        target_part = getattr(part, "related_parts", {}).get(rid)
        if target_part is None and hasattr(part, "rels") and rid in part.rels:
            target_part = part.rels[rid].target_part
        if target_part is None or not hasattr(target_part, "blob"):
            continue
        blob = target_part.blob
        sha = hashlib.sha256(blob).hexdigest()
        exts = drawing_elem.xpath(".//wp:extent")
        cx = int(exts[0].attrib.get("cx")) if exts and exts[0].attrib.get("cx") else None
        cy = int(exts[0].attrib.get("cy")) if exts and exts[0].attrib.get("cy") else None
        doc_prs = drawing_elem.xpath(".//wp:docPr")
        descr = (
            (doc_prs[0].attrib.get("descr") or doc_prs[0].attrib.get("name") or "")
            if doc_prs
            else ""
        )
        if sha in seen_hashes:
            return seen_hashes[sha], cx, cy, descr

        key = f"img_{sha[:12]}"
        seen_hashes[sha] = key
        file_path_str = None
        if cache_dir is not None:
            ext = (
                ".png"
                if "png" in target_part.content_type
                else ".jpg"
                if "jpeg" in target_part.content_type
                else ".bin"
            )
            asset_path = cache_dir / f"{key}{ext}"
            try:
                cache_dir.mkdir(parents=True, exist_ok=True)
                asset_path.write_bytes(blob)
                file_path_str = str(asset_path)
            except OSError:
                file_path_str = None

        is_logo = (
            location in ("header", "footer")
            or "logo" in descr.lower()
            or "brand" in descr.lower()
            or (location == "body" and section_index == 0 and cx is not None and cx < 3000000)
        )
        image_assets.append(
            ImageAssetInfo(
                asset_key=key,
                original_rel_id=rid,
                content_type=target_part.content_type,
                location=location,
                section_index=section_index,
                width_emu=cx,
                height_emu=cy,
                description=descr,
                is_logo=is_logo,
                blob_sha256=sha,
                file_path=file_path_str,
            )
        )
        return key, cx, cy, descr
    return None, None, None, ""


def _inspect_header_footer_part(
    hf_part,
    location: str,
    section_index: int,
    image_assets: list[ImageAssetInfo],
    seen_hashes: dict[str, str],
    cache_dir: Path | None,
) -> HeaderFooterInfo:
    if hf_part is None:
        return HeaderFooterInfo(present=False)

    elements: list[HeaderFooterElementInfo] = []
    has_image = False
    has_table = len(hf_part.tables) > 0

    for p in hf_part.paragraphs:
        # Check drawings in paragraph
        drawings = p._p.xpath(".//w:drawing | .//w:pict")
        for drawing in drawings:
            key, cx, cy, descr = _extract_image_asset(
                hf_part.part, drawing, location, section_index, image_assets, seen_hashes, cache_dir
            )
            if key:
                has_image = True
                elements.append(
                    HeaderFooterElementInfo(
                        kind="image",
                        image_asset_key=key,
                        width_emu=cx,
                        height_emu=cy,
                        text=descr or None,
                    )
                )

        # Check fields in paragraph (e.g. PAGE, NUMPAGES, DATE)
        field_codes: list[str] = []
        for instr in p._p.xpath(".//w:instrText"):
            if instr.text:
                field_codes.append(instr.text.strip())
        for fld in p._p.xpath(".//w:fldSimple"):
            instr = fld.attrib.get(qn("w:instr")) or fld.attrib.get("instr")
            if instr:
                field_codes.append(instr.strip())

        for fc in field_codes:
            code_word = fc.split()[0].upper() if fc.split() else ""
            if code_word in ("PAGE", "NUMPAGES", "DATE", "TIME", "FILENAME", "TITLE", "AUTHOR"):
                elements.append(HeaderFooterElementInfo(kind="field", field_code=code_word, text=fc))

        if p.text.strip():
            elements.append(HeaderFooterElementInfo(kind="text", text=p.text))

    for tbl in hf_part.tables:
        tbl_text = "\n".join(
            " | ".join(c.text.strip() for c in r.cells if c.text.strip())
            for r in tbl.rows
        ).strip()
        elements.append(HeaderFooterElementInfo(kind="table", text=tbl_text or None))
        for row in tbl.rows:
            for cell in row.cells:
                for drawing in cell._tc.xpath(".//w:drawing | .//w:pict"):
                    key, cx, cy, descr = _extract_image_asset(
                        hf_part.part, drawing, location, section_index, image_assets, seen_hashes, cache_dir
                    )
                    if key:
                        has_image = True
                        elements.append(
                            HeaderFooterElementInfo(
                                kind="image",
                                image_asset_key=key,
                                width_emu=cx,
                                height_emu=cy,
                                text=descr or None,
                            )
                        )

    full_text = "\n".join(p.text for p in hf_part.paragraphs if p.text.strip())
    present = bool(full_text or elements or has_table or has_image)
    return HeaderFooterInfo(
        present=present,
        text=full_text,
        elements=elements,
        has_table=has_table,
        has_image=has_image,
    )


def analyze_docx(path: Path, *, template_id: str | None = None) -> TemplateManifest:
    data = path.read_bytes()
    doc = Document(str(path))
    used = {paragraph.style.style_id for paragraph in doc.paragraphs if paragraph.style is not None}
    used.update(table.style.style_id for table in doc.tables if table.style is not None)
    styles = [
        StyleInfo(
            name=style.name,
            style_id=style.style_id,
            type=style.type.__str__(),
            based_on=getattr(getattr(style, "base_style", None), "style_id", None),
            font=_font(style),
            in_use=style.style_id in used,
        )
        for style in doc.styles
    ]

    theme_colors = _extract_theme_colors(path)
    image_assets: list[ImageAssetInfo] = []
    seen_hashes: dict[str, str] = {}
    cache_dir = path.parent / f".{path.stem}-assets"

    sections: list[SectionInfo] = []
    for s_idx, section in enumerate(doc.sections):
        default_hdr = _inspect_header_footer_part(
            section.header, "header", s_idx, image_assets, seen_hashes, cache_dir
        )
        default_ftr = _inspect_header_footer_part(
            section.footer, "footer", s_idx, image_assets, seen_hashes, cache_dir
        )
        first_hdr = _inspect_header_footer_part(
            getattr(section, "first_page_header", None), "header", s_idx, image_assets, seen_hashes, cache_dir
        )
        first_ftr = _inspect_header_footer_part(
            getattr(section, "first_page_footer", None), "footer", s_idx, image_assets, seen_hashes, cache_dir
        )
        even_hdr = _inspect_header_footer_part(
            getattr(section, "even_page_header", None), "header", s_idx, image_assets, seen_hashes, cache_dir
        )
        even_ftr = _inspect_header_footer_part(
            getattr(section, "even_page_footer", None), "footer", s_idx, image_assets, seen_hashes, cache_dir
        )

        has_pg = (
            "PAGE" in section.header._element.xml
            or "PAGE" in section.footer._element.xml
            or any(el.field_code == "PAGE" for el in default_ftr.elements + first_ftr.elements)
        )
        sections.append(
            SectionInfo(
                page_width=int(section.page_width),
                page_height=int(section.page_height),
                orientation=section.orientation.name.lower(),
                margins={
                    "top": int(section.top_margin),
                    "bottom": int(section.bottom_margin),
                    "left": int(section.left_margin),
                    "right": int(section.right_margin),
                },
                header_text=default_hdr.text,
                footer_text=default_ftr.text,
                has_page_number=has_pg,
                different_first_page=bool(section.different_first_page_header_footer),
                default_header=default_hdr,
                default_footer=default_ftr,
                first_page_header=first_hdr,
                first_page_footer=first_ftr,
                even_page_header=even_hdr,
                even_page_footer=even_ftr,
                header_distance_emu=int(section.header_distance) if section.header_distance is not None else None,
                footer_distance_emu=int(section.footer_distance) if section.footer_distance is not None else None,
            )
        )

    # Also scan body for cover page / branding images
    for p in doc.paragraphs[:10]:
        for drawing in p._p.xpath(".//w:drawing | .//w:pict"):
            _extract_image_asset(
                doc.part, drawing, "body", 0, image_assets, seen_hashes, cache_dir
            )

    by_name = {style.name.lower(): style.name for style in doc.styles}
    role_map = {
        role: by_name[name]
        for role, name in (
            ("title", "title"),
            ("h1", "heading 1"),
            ("h2", "heading 2"),
            ("h3", "heading 3"),
            ("body", "normal"),
            ("caption", "caption"),
        )
        if name in by_name
    }
    normal = next((style for style in styles if style.name.lower() == "normal"), None)
    heading = next((style for style in styles if style.name.lower() == "heading 1"), None)
    return TemplateManifest(
        id=template_id or f"tpl_{hashlib.sha256(data).hexdigest()[:12]}",
        format=OutputFormat.docx,
        file_sha256=hashlib.sha256(data).hexdigest(),
        theme=ThemeSpec(
            major_font=heading.font.name if heading and heading.font else None,
            minor_font=normal.font.name if normal and normal.font else None,
            colors=theme_colors,
        ),
        docx=DocxTemplateInfo(
            styles=styles,
            sections=sections,
            has_toc="TOC" in doc.element.body.xml,
            paragraph_count=len(doc.paragraphs),
            table_count=len(doc.tables),
            role_style_map=role_map,
            image_assets=image_assets,
            theme_colors=theme_colors,
        ),
        analyzer_version=ANALYZER_VERSION,
    )
