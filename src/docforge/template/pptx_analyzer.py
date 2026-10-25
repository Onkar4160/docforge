"""PPTX template analysis → TemplateManifest (Stage 6, FR-TPL-2/3/4/7).

Reads the theme (colours, fonts), every layout's placeholders, and every element on the sample slides,
classifying elements as structural, branding, decorative, placeholder or sample content.
"""

from __future__ import annotations

import hashlib
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from docforge.models.intent import OutputFormat
from docforge.models.template import (
    Box,
    ElementKind,
    ElementRole,
    ElementSpec,
    FontSpec,
    LayoutArchetype,
    LayoutSpec,
    PlaceholderSpec,
    PptxTemplateInfo,
    PropertySource,
    SlideInventory,
    TemplateManifest,
    ThemeSpec,
)
from docforge.template.slide_design import analyse_slide, design_mode, slide_style

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_SLOTS = (
    "dk1",
    "lt1",
    "dk2",
    "lt2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
)
_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_BODY_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}
_META_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
_POS_TOLERANCE = int(0.1 * 914400)
# Bump when analysis output changes so cached manifests are recomputed (2: designed-slide bases).
ANALYZER_VERSION = 2


def _theme(prs) -> ThemeSpec:
    try:
        theme_part = prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme_part.blob)
    except (KeyError, etree.XMLSyntaxError):
        return ThemeSpec()
    colors: dict[str, str] = {}
    scheme = root.find(f".//{_A}clrScheme")
    if scheme is not None:
        for slot in _SLOTS:
            node = scheme.find(f"{_A}{slot}")
            if node is None or not len(node):
                continue
            value = node[0].get("val") if node[0].tag == f"{_A}srgbClr" else node[0].get("lastClr")
            if value:
                colors[slot] = value.upper()
    major = root.find(f".//{_A}majorFont/{_A}latin")
    minor = root.find(f".//{_A}minorFont/{_A}latin")
    return ThemeSpec(
        colors=colors,
        major_font=major.get("typeface") if major is not None else None,
        minor_font=minor.get("typeface") if minor is not None else None,
    )


def _box(shape) -> Box | None:
    if shape.left is None or shape.width is None:
        return None
    return Box(x=int(shape.left), y=int(shape.top), w=int(shape.width), h=int(shape.height))


def _kind(shape) -> ElementKind:
    if shape.is_placeholder:
        return ElementKind.placeholder
    t = shape.shape_type
    if t == MSO_SHAPE_TYPE.PICTURE:
        return ElementKind.picture
    if t == MSO_SHAPE_TYPE.GROUP:
        return ElementKind.group
    if t == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False):
        return ElementKind.table
    if getattr(shape, "has_chart", False):
        return ElementKind.chart
    if t == MSO_SHAPE_TYPE.LINE or shape.__class__.__name__ == "Connector":
        return ElementKind.line
    if shape.has_text_frame and shape.text_frame.text.strip():
        return ElementKind.text
    if t in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
        return ElementKind.shape
    return ElementKind.other


def _font(shape) -> FontSpec | None:
    if not shape.has_text_frame:
        return None
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            color = None
            try:
                if r.font.color is not None and r.font.color.type is not None:
                    try:
                        color = str(r.font.color.rgb)
                    except AttributeError:
                        theme_color = r.font.color.theme_color
                        if theme_color is not None:
                            color = f"theme:{theme_color.name.lower().replace('_', '')}"
            except AttributeError:
                color = None

            def source(value: object, *, has_explicit_color: bool = False) -> PropertySource:
                if value is not None or has_explicit_color:
                    return PropertySource.explicit
                return PropertySource.inherited if shape.is_placeholder else PropertySource.unknown

            return FontSpec(
                name=r.font.name,
                size_pt=r.font.size.pt if r.font.size else None,
                bold=r.font.bold,
                color=color,
                name_source=source(r.font.name),
                size_source=source(r.font.size),
                bold_source=source(r.font.bold),
                color_source=source(
                    color,
                    has_explicit_color=r.font.color is not None and r.font.color.type is not None,
                ),
            )
    return None


def _fill(shape) -> tuple[str | None, PropertySource]:
    """Return an explicit RGB fill when available; unresolved/inherited fills keep their provenance."""
    try:
        fill = shape.fill
        if fill.type is None:
            return None, PropertySource.inherited if shape.is_placeholder else PropertySource.unknown
        try:
            color = str(fill.fore_color.rgb).upper()
        except (AttributeError, TypeError, ValueError):
            try:
                theme_color = fill.fore_color.theme_color
            except (AttributeError, TypeError, ValueError):
                theme_color = None
            if theme_color is not None:
                return f"theme:{theme_color.name.lower().replace('_', '')}", PropertySource.explicit
            return None, PropertySource.unknown
        return color, PropertySource.explicit
    except (AttributeError, TypeError, ValueError):
        return None, PropertySource.unknown


def _placeholder_box(ph, layout) -> tuple[Box | None, PropertySource]:
    """Placeholders may inherit geometry from the master; resolve it."""
    box = _box(ph)
    if box is not None:
        return box, PropertySource.explicit
    ptype = ph.placeholder_format.type
    for candidate in layout.slide_master.placeholders:
        if candidate.placeholder_format.type == ptype or (
            ptype in _TITLE_TYPES and candidate.placeholder_format.type in _TITLE_TYPES
        ):
            inherited = _box(candidate)
            if inherited is not None:
                return inherited, PropertySource.inherited
    return None, PropertySource.unknown


def classify_layout(name: str, placeholders: list[PlaceholderSpec]) -> LayoutArchetype:
    n = name.lower()
    types = Counter(p.type for p in placeholders)
    titles = types["center_title"] + types["title"]
    bodies = types["body"] + types["object"]
    pictures = types["picture"]
    if any(k in n for k in ("closing", "thank", "end slide", "conclusion")):
        return LayoutArchetype.closing
    if types["center_title"] or "title slide" in n or n.strip() in ("title", "cover"):
        return LayoutArchetype.title
    if "section" in n or "divider" in n:
        return LayoutArchetype.section
    if "comparison" in n:
        return LayoutArchetype.comparison
    if pictures and titles:
        return LayoutArchetype.image_text
    if titles and bodies >= 2:
        return LayoutArchetype.two_column
    if titles and bodies == 1:
        return LayoutArchetype.content
    if titles and not bodies:
        return LayoutArchetype.title_only
    if not titles and not bodies:
        return LayoutArchetype.blank
    return LayoutArchetype.other


def _ptype_name(ph) -> str:
    try:
        return ph.placeholder_format.type.name.lower() if ph.placeholder_format.type else "body"
    except AttributeError:
        return "body"


def analyze_pptx(path: Path, *, template_id: str | None = None) -> TemplateManifest:
    data = path.read_bytes()
    prs = Presentation(str(path))
    theme = _theme(prs)
    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    warnings: list[str] = []

    layouts: list[LayoutSpec] = []
    layout_ids: dict[int, str] = {}
    for m_i, master in enumerate(prs.slide_masters):
        for l_i, layout in enumerate(master.slide_layouts):
            lid = f"m{m_i}l{l_i}"
            layout_ids[id(layout)] = lid
            placeholders = []
            for ph in layout.placeholders:
                if ph.placeholder_format.type in _META_TYPES:
                    continue
                bbox, bbox_source = _placeholder_box(ph, layout)
                placeholders.append(
                    PlaceholderSpec(
                        idx=ph.placeholder_format.idx,
                        type=_ptype_name(ph),
                        name=ph.name,
                        bbox=bbox,
                        bbox_source=bbox_source,
                        font=_font(ph),
                    )
                )
            elements = []
            for n, shape in enumerate(layout.shapes):
                if shape.is_placeholder:
                    continue
                fill, fill_source = _fill(shape)
                elements.append(
                    ElementSpec(
                        id=f"{lid}-e{n}",
                        name=shape.name,
                        kind=_kind(shape),
                        bbox=_box(shape),
                        bbox_source=PropertySource.explicit,
                        z=n,
                        role=(
                            ElementRole.branding
                            if _kind(shape) == ElementKind.picture
                            else ElementRole.structural
                        ),
                        fill=fill,
                        fill_source=fill_source,
                        font=_font(shape),
                    )
                )
            layouts.append(
                LayoutSpec(
                    id=lid,
                    name=layout.name,
                    master_index=m_i,
                    archetype=classify_layout(layout.name, placeholders),
                    placeholders=placeholders,
                    elements=elements,
                )
            )

    # Sample slides: inventory + role classification.
    inventories: list[SlideInventory] = []
    pictures: Counter = Counter()
    texts: Counter = Counter()
    slides = list(prs.slides)
    for slide in slides:
        for s in slide.shapes:
            box = _box(s)
            if _kind(s) == ElementKind.picture and box:
                pictures[(box.x // _POS_TOLERANCE, box.y // _POS_TOLERANCE, box.w // _POS_TOLERANCE)] += 1
            if _kind(s) == ElementKind.text and s.has_text_frame:
                texts[s.text_frame.text.strip()] += 1
    repeated = max(2, (len(slides) + 1) // 2)
    for i, slide in enumerate(slides):
        elements = []
        for n, s in enumerate(slide.shapes):
            kind, box = _kind(s), _box(s)
            fill, fill_source = _fill(s)
            role = ElementRole.sample_content
            if kind == ElementKind.placeholder:
                has_text = s.has_text_frame and s.text_frame.text.strip()
                role = ElementRole.sample_content if has_text else ElementRole.placeholder
            elif kind == ElementKind.picture and box:
                key = (box.x // _POS_TOLERANCE, box.y // _POS_TOLERANCE, box.w // _POS_TOLERANCE)
                small = box.area() < 0.03 * slide_w * slide_h
                role = (
                    ElementRole.branding if pictures[key] >= repeated or small else ElementRole.sample_content
                )
            elif (
                kind == ElementKind.text
                and s.has_text_frame
                and texts[s.text_frame.text.strip()] >= repeated
                and len(slides) > 1
            ):
                role = ElementRole.branding
            elif kind in (ElementKind.shape, ElementKind.line):
                role = ElementRole.decorative
            text = s.text_frame.text.strip()[:300] if s.has_text_frame and s.text_frame.text.strip() else None
            ph_type = _ptype_name(s) if s.is_placeholder else None
            elements.append(
                ElementSpec(
                    id=f"s{i}-e{n}",
                    name=s.name,
                    kind=kind,
                    bbox=box,
                    z=n,
                    role=role,
                    text=text,
                    fill=fill,
                    fill_source=fill_source,
                    bbox_source=PropertySource.explicit if box is not None else PropertySource.unknown,
                    font=_font(s),
                    placeholder_type=ph_type,
                    placeholder_idx=s.placeholder_format.idx if s.is_placeholder else None,
                )
            )
        inventories.append(
            SlideInventory(index=i, layout_id=layout_ids.get(id(slide.slide_layout), "?"), elements=elements)
        )

    bases = []
    for i, slide in enumerate(slides):
        try:
            base = analyse_slide(i, slide, slide_w, slide_h)
        except Exception as exc:  # noqa: BLE001 - one odd slide must not break template analysis
            warnings.append(f"sample slide {i + 1} could not be analysed ({exc.__class__.__name__})")
            continue
        if base is not None:
            bases.append(base)
    mode = design_mode(prs, bases)
    fonts, palette, text_color, background = ({}, [], None, None)
    if mode == "slides":
        fonts, palette, text_color, background = slide_style(prs, bases, theme.colors)

    if mode == "layouts" and not any(
        lay.archetype in (LayoutArchetype.title_only, LayoutArchetype.content) for lay in layouts
    ):
        warnings.append(
            "template has no title-only or title+content layout; content slides use the closest layout"
        )
    if not theme.colors:
        warnings.append("theme colours could not be read; generated palette used")
    return TemplateManifest(
        id=template_id or f"tpl_{hashlib.sha256(data).hexdigest()[:12]}",
        format=OutputFormat.pptx,
        file_sha256=hashlib.sha256(data).hexdigest(),
        theme=theme,
        pptx=PptxTemplateInfo(
            slide_width=slide_w,
            slide_height=slide_h,
            master_count=len(prs.slide_masters),
            layouts=layouts,
            sample_slides=inventories,
            design_mode=mode,
            slide_bases=bases,
            slide_fonts=fonts,
            slide_palette=palette,
            slide_text_color=text_color,
            slide_background=background,
        ),
        warnings=warnings,
        analyzer_version=ANALYZER_VERSION,
    )
