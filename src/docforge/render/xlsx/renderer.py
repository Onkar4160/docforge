"""WorkbookPlan → editable XLSX with real formulas (Stage 10, FR-XLSX-1..5, 7).

Adds a README sheet (purpose, sheet guide, colour conventions) and, when evidence is cited, a Sources sheet.
Inputs use the modelling convention: blue text on a light fill; outputs are bold.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from docforge.design.system import readable_on
from docforge.ingest.datasets import Dataset
from docforge.models.design import DesignTokens
from docforge.models.evidence import ResearchDossier
from docforge.models.ir import ChartType, ColumnType, SheetPlan, TableSpec, WorkbookPlan
from docforge.render.citations import Citations

_FORMATS = {
    ColumnType.int: "#,##0",
    ColumnType.float: "#,##0.00",
    ColumnType.percent: "0.0%",
    ColumnType.date: "yyyy-mm-dd",
}
_THIN = Side(style="thin", color="D9DDE3")


@dataclass
class WorkbookRenderResult:
    path: Path
    sheet_names: dict[str, str] = field(default_factory=dict)  # plan sheet id -> sheet name
    table_ranges: dict[str, str] = field(default_factory=dict)  # "Sheet!A1:D20" per table
    calculated_cells: list[str] = field(default_factory=list)  # "Sheet!E2" cells that must hold formulas
    typed_cells: list[tuple[str, str]] = field(default_factory=list)  # (cell, expected number format)
    sources: int = 0


def number_format(dtype: ColumnType, unit: str | None, explicit: str | None) -> str | None:
    if explicit:
        return explicit
    if dtype == ColumnType.currency:
        symbol = (unit or "").strip()
        if symbol in ("$", "USD"):
            return '"$"#,##0.00'
        if symbol in ("€", "EUR"):
            return '#,##0.00 "€"'
        if symbol in ("£", "GBP"):
            return '"£"#,##0.00'
        return f'#,##0.00 "{symbol}"' if symbol else "#,##0.00"
    return _FORMATS.get(dtype)


def _table_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "_", name) or "Table"
    if not base[0].isalpha():
        base = f"T_{base}"
    candidate, n = base[:200], 2
    while candidate.lower() in used:
        candidate, n = f"{base[:190]}_{n}", n + 1
    used.add(candidate.lower())
    return candidate


class XlsxRenderer:
    def __init__(self, tokens: DesignTokens, datasets: dict[str, Dataset], citations: Citations) -> None:
        self.t = tokens
        self.datasets = datasets
        self.cite = citations
        self.header_fill = PatternFill("solid", fgColor=tokens.palette.primary)
        self.header_font = Font(bold=True, color=readable_on(tokens.palette.primary))
        self.input_fill = PatternFill("solid", fgColor="FFF7E0")
        self.input_font = Font(color="1F4FD1")
        self._tables: set[str] = set()

    def render(self, plan: WorkbookPlan, path: Path) -> WorkbookRenderResult:
        wb = Workbook()
        wb.remove(wb.active)
        result = WorkbookRenderResult(path=path)
        result.sheet_names = {s.id: s.name for s in plan.sheets}
        readme = wb.create_sheet("README")
        for sheet in plan.sheets:
            self._sheet(wb.create_sheet(sheet.name), sheet, result)
        for nr in plan.named_ranges:
            sheet_name = result.sheet_names.get(nr.sheet_id)
            if not sheet_name:
                continue
            col, row = coordinate_from_string(nr.ref.replace("$", "").split(":")[0])
            wb.defined_names[nr.name] = DefinedName(nr.name, attr_text=f"'{sheet_name}'!${col}${row}")
        cited = [cid for s in plan.sheets for cid in s.claim_ids]
        if cited:
            self.cite.marker(cited)
            entries = self.cite.reference_entries()
            if entries:
                self._sources(wb.create_sheet("Sources"), entries)
                result.sources = len(entries)
        self._readme(readme, plan, result)
        wb.properties.title = plan.title[:255]
        wb.properties.creator = "DocForge"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return result

    # ---------------------------------------------------------------- sheets

    def _sheet(self, ws, sheet: SheetPlan, result: WorkbookRenderResult) -> None:
        widths: dict[int, int] = {}
        for table in sheet.tables:
            self._table(ws, sheet, table, result, widths)
        for cell in sheet.cells:
            target = ws[cell.ref]
            target.value = cell.formula if cell.formula else cell.value
            if cell.number_format:
                target.number_format = cell.number_format
            role = cell.style_role
            if role == "input":
                target.fill, target.font = self.input_fill, self.input_font
            elif role == "output":
                target.font = Font(bold=True)
            elif role == "header":
                target.fill, target.font = self.header_fill, self.header_font
            elif role == "label":
                target.font = Font(bold=True, color=self.t.palette.text)
            elif role == "note":
                target.font = Font(italic=True, color=self.t.palette.text_muted)
            col = target.column
            widths[col] = max(widths.get(col, 8), min(60, len(str(target.value or "")) + 2))
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        if sheet.freeze:
            ws.freeze_panes = sheet.freeze
        for spec in sheet.charts:
            self._chart(ws, spec)
        ws.sheet_view.showGridLines = sheet.role.value not in ("report",)

    def _table(self, ws, sheet: SheetPlan, table: TableSpec, result, widths: dict[int, int]) -> None:
        col0_letter, row0 = coordinate_from_string(table.anchor)
        col0 = column_index_from_string(col0_letter)
        rows = table.rows
        if table.source_dataset and table.source_dataset in self.datasets:
            ds = self.datasets[table.source_dataset]
            index = {c.name.lower(): i for i, c in enumerate(ds.columns)}
            rows = [
                [r[index[c.name.lower()]] if c.name.lower() in index else None for c in table.columns]
                for r in ds.rows
            ]
        for j, column in enumerate(table.columns):
            cell = ws.cell(
                row=row0,
                column=col0 + j,
                value=column.name
                + (f" ({column.unit})" if column.unit and column.dtype != ColumnType.currency else ""),
            )
            cell.fill, cell.font = self.header_fill, self.header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            widths[col0 + j] = max(widths.get(col0 + j, 8), min(40, len(str(cell.value)) + 3))
        for i, values in enumerate(rows, start=1):
            r = row0 + i
            for j, column in enumerate(table.columns):
                value = values[j] if j < len(values) else None
                cell = ws.cell(row=r, column=col0 + j)
                if column.formula_template:
                    cell.value = column.formula_template.format(row=r)
                    result.calculated_cells.append(f"{sheet.name}!{cell.coordinate}")
                else:
                    cell.value = value
                fmt = number_format(column.dtype, column.unit, column.number_format)
                if fmt:
                    cell.number_format = fmt
                    result.typed_cells.append((f"{sheet.name}!{cell.coordinate}", fmt))
                cell.border = Border(bottom=_THIN)
                if value is not None and not column.formula_template:
                    widths[col0 + j] = max(widths.get(col0 + j, 8), min(40, len(str(value)) + 2))
        last_col = get_column_letter(col0 + len(table.columns) - 1)
        last_row = row0 + max(len(rows), 1)
        ref = f"{col0_letter}{row0}:{last_col}{last_row}"
        result.table_ranges[f"{sheet.name}!{table.name}"] = f"{sheet.name}!{ref}"
        if rows:
            xl_table = Table(displayName=_table_name(table.name, self._tables), ref=ref)
            xl_table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
            ws.add_table(xl_table)

    def _chart(self, ws, spec) -> None:
        kind = spec.chart_type
        chart = {
            ChartType.line: LineChart,
            ChartType.area: AreaChart,
            ChartType.pie: PieChart,
            ChartType.doughnut: PieChart,
        }.get(kind, BarChart)()
        if isinstance(chart, BarChart):
            chart.type = "bar" if kind in (ChartType.bar, ChartType.stacked_bar) else "col"
            if kind in (ChartType.stacked_bar, ChartType.stacked_column):
                chart.grouping, chart.overlap = "stacked", 100
        min_col, min_row, max_col, max_row = range_boundaries(spec.data_ref)
        data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        c_min_col, c_min_row, c_max_col, c_max_row = range_boundaries(spec.categories_ref)
        cats = Reference(ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.title = spec.title
        chart.height, chart.width = 8, 16
        if not isinstance(chart, PieChart):
            for i, series in enumerate(chart.series):
                color = self.t.palette.accents[i % max(1, len(self.t.palette.accents) - 1)]
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.solidFill = color
        ws.add_chart(chart, spec.anchor)

    # ---------------------------------------------------------------- generated sheets

    def _readme(self, ws, plan: WorkbookPlan, result: WorkbookRenderResult) -> None:
        ws["A1"] = plan.title
        ws["A1"].font = Font(bold=True, size=16, color=self.t.palette.primary)
        ws["A2"] = f"Generated by DocForge on {datetime.now(UTC):%Y-%m-%d}."
        ws["A2"].font = Font(italic=True, color=self.t.palette.text_muted)
        ws["A4"], ws["B4"], ws["C4"] = "Sheet", "Role", "What it contains"
        for c in ("A4", "B4", "C4"):
            ws[c].fill, ws[c].font = self.header_fill, self.header_font
        row = 5
        for sheet in plan.sheets:
            ws.cell(row=row, column=1, value=sheet.name)
            ws.cell(row=row, column=2, value=sheet.role.value)
            ws.cell(row=row, column=3, value=sheet.description)
            row += 1
        if result.sources:
            ws.cell(row=row, column=1, value="Sources")
            ws.cell(row=row, column=2, value="sources")
            ws.cell(row=row, column=3, value="References for figures taken from research or your files")
            row += 1
        row += 1
        ws.cell(row=row, column=1, value="Conventions").font = Font(bold=True)
        legend = ws.cell(row=row + 1, column=1, value="Input (edit me)")
        legend.fill, legend.font = self.input_fill, self.input_font
        ws.cell(row=row + 1, column=2, value="Blue text on a light fill: assumptions you can change.")
        ws.cell(row=row + 2, column=1, value="Output").font = Font(bold=True)
        ws.cell(
            row=row + 2, column=2, value="Bold: calculated results. All calculations are live Excel formulas."
        )
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 80
        ws.sheet_view.showGridLines = False

    def _sources(self, ws, entries: list[str]) -> None:
        ws["A1"] = "Sources"
        ws["A1"].font = Font(bold=True, size=14, color=self.t.palette.primary)
        for i, entry in enumerate(entries, start=3):
            ws.cell(row=i, column=1, value=entry).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 120


def render_workbook(
    plan: WorkbookPlan,
    tokens: DesignTokens,
    out_path: Path,
    *,
    datasets: dict[str, Dataset],
    dossier: ResearchDossier | None = None,
    citation_style="numeric",
) -> WorkbookRenderResult:
    return XlsxRenderer(tokens, datasets, Citations(dossier, citation_style)).render(plan, out_path)
